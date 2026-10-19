import re
from html import escape
from pathlib import Path
from typing import Any
from urllib.parse import quote
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


REPORTS_DIR = Path(__file__).resolve().parent / "generated_reports"
PROJECT_DIR = Path(__file__).resolve().parent.parent
_SECRET_PATTERN = re.compile(
    r"(?i)\b(password|passwd|token|api[_-]?key|secret|authorization)"
    r"\b\s*[:=]\s*(['\"]?)([^\s,'\";]+)"
)
_BEARER_PATTERN = re.compile(r"(?i)\bBearer\s+[A-Za-z0-9._~+/=-]+")
_ASSERTION_ACTIONS = {
    "check_visibility",
    "verify_authentication",
    "autonomous:assert_visible",
    "autonomous:assert_outcome",
}


def _assertion_outcome(step: Any) -> str:
    if step.action not in _ASSERTION_ACTIONS:
        return "Not evaluated"
    if step.status == "passed":
        return "Passed"
    if step.status == "failed":
        return "Failed"
    return "Not evaluated"


def _text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = _BEARER_PATTERN.sub("Bearer [REDACTED]", text)
    text = _SECRET_PATTERN.sub(r"\1=[REDACTED]", text)
    return text


def _report_directory(run_id: str) -> Path:
    safe_id = str(UUID(run_id))
    output = REPORTS_DIR / safe_id
    output.mkdir(parents=True, exist_ok=True)
    resolved = output.resolve()
    if not resolved.is_relative_to(REPORTS_DIR.resolve()):
        raise ValueError("Report output directory escaped its configured root.")
    return resolved


def resolve_evidence_path(run_id: str, stored_path: str | None) -> Path | None:
    if not stored_path:
        return None
    safe_id = str(UUID(run_id))
    relative = Path(stored_path.replace("\\", "/"))
    expected_root = Path("evidence") / safe_id
    if relative.is_absolute() or ".." in relative.parts:
        return None
    if len(relative.parts) < 3 or relative.parts[:2] != expected_root.parts:
        return None

    evidence_root = PROJECT_DIR / expected_root
    resolved_root = evidence_root.resolve()
    resolved_file = (PROJECT_DIR / relative).resolve()
    if not resolved_file.is_relative_to(resolved_root) or not resolved_file.is_file():
        return None
    return resolved_file


def generate_excel_report(data: dict[str, Any]) -> Path:
    run = data["run"]
    run_id = str(run.id)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Test Summary"
    summary.append(["Field", "Value"])
    for field, value in (
        ("Run ID", run_id),
        ("Test name", run.run_name),
        ("Platform", run.platform),
        ("Mode", run.mode),
        ("Status", run.status),
        ("Target", run.target),
        ("Assertion result", run.assertion_result),
        ("Suspected bug count", run.suspected_bug_count),
        ("Assertions passed", run.assertion_passed_count),
        ("Assertions failed", run.assertion_failed_count),
        ("Execution error", run.execution_error),
        ("Created at", run.created_at),
        ("Started at", run.started_at),
        ("Ended at", run.ended_at),
        ("Duration (ms)", _duration_ms(run.started_at, run.ended_at)),
        ("Failure summary", _failure_summary(run, data["steps"])),
        ("Goal", run.goal),
        ("Expected results", "\n".join(run.expected_results)),
    ):
        summary.append([field, _text(value)])

    steps_sheet = workbook.create_sheet("Test Steps")
    steps_sheet.append(
        [
            "Step number",
            "Action",
            "Page or screen",
            "Expected result",
            "Actual result",
            "Status",
            "Assertion outcome",
            "Duration (ms)",
            "Error",
            "Evidence availability",
            "Evidence reference",
            "Timestamp",
        ]
    )
    for step in data["steps"]:
        evidence = resolve_evidence_path(run_id, step.screenshot_path)
        steps_sheet.append(
            [
                step.step_number,
                _text(step.action),
                _text(step.page_or_screen),
                _text(step.expected_result),
                _text(step.actual_result),
                step.status,
                (
                    _assertion_outcome(step)
                ),
                step.duration_ms,
                _text(step.error_details),
                "Available" if evidence else "Unavailable",
                str(evidence.relative_to(PROJECT_DIR)) if evidence else "",
                step.created_at.isoformat(),
            ]
        )

    bug_sheet = workbook.create_sheet("Bug Reports")
    bug_sheet.append(
        [
            "Title",
            "Description",
            "Severity",
            "Confidence",
            "Reproduction steps",
            "Expected result",
            "Actual result",
            "Evidence",
            "Relevant logs",
            "Created at",
        ]
    )
    for bug in data["bugs"]:
        evidence = resolve_evidence_path(run_id, bug.evidence_path)
        bug_sheet.append(
            [
                _text(bug.title),
                _text(bug.description),
                bug.severity,
                bug.confidence,
                "\n".join(_text(item) for item in bug.reproduction_steps),
                _text(bug.expected_result),
                _text(bug.actual_result),
                str(evidence.relative_to(PROJECT_DIR)) if evidence else "",
                "\n".join(_text(item) for item in bug.relevant_logs),
                bug.created_at.isoformat(),
            ]
        )

    log_sheet = workbook.create_sheet("Execution Logs")
    log_sheet.append(["Timestamp", "Event", "Message", "Step ID"])
    for log in data["logs"]:
        log_sheet.append(
            [
                log.timestamp.isoformat(),
                _text(log.event_type),
                _text(log.message),
                log.step_id or "",
            ]
        )

    browser_errors = data.get("browser_errors") or []
    if browser_errors:
        browser_sheet = workbook.create_sheet("Browser Errors")
        browser_sheet.append(
            [
                "Step number",
                "Category",
                "HTTP status",
                "Resource type",
                "URL",
                "Message",
                "Timestamp",
            ]
        )
        for error in browser_errors:
            browser_sheet.append(
                [
                    error.step_number if error.step_number is not None else "",
                    error.category,
                    error.http_status if error.http_status is not None else "",
                    error.resource_type or "",
                    _text(error.url),
                    _text(error.message),
                    error.timestamp.isoformat(),
                ]
            )

    ai_reports = data.get("ai_bug_reports") or []
    if ai_reports:
        ai_sheet = workbook.create_sheet("AI Failure Analysis")
        ai_sheet.append(["Title", "Model", "Created at", "Analysis Report"])
        for rep in ai_reports:
            ai_sheet.append(
                [
                    _text(rep.title),
                    _text(rep.model),
                    rep.created_at.isoformat(),
                    _text(rep.report_markdown),
                ]
            )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
        for column in sheet.columns:
            width = min(
                60,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )
            sheet.column_dimensions[column[0].column_letter].width = width

    output_path = _report_directory(run_id) / "report.xlsx"
    workbook.save(output_path)
    return output_path


def generate_html_report(data: dict[str, Any], base_url: str) -> Path:
    run = data["run"]
    run_id = str(run.id)
    rows = []
    for step in data["steps"]:
        screenshot = resolve_evidence_path(run_id, step.screenshot_path)
        screenshot_html = "No screenshot evidence recorded."
        if screenshot:
            href = (
                f"{base_url.rstrip('/')}/test-runs/{quote(run_id)}/steps/"
                f"{step.step_number}/screenshot"
            )
            screenshot_html = (
                f'<a href="{escape(href, quote=True)}">View screenshot</a>'
                f' <span>({escape(screenshot.name)})</span>'
            )
        rows.append(
            "<article class=\"step\">"
            f"<h3>Step {step.step_number}: {escape(_text(step.action))}</h3>"
            f"<p><strong>Status:</strong> {escape(step.status)}</p>"
            f"<p><strong>Expected:</strong> {escape(_text(step.expected_result)) or 'Not recorded'}</p>"
            f"<p><strong>Actual:</strong> {escape(_text(step.actual_result)) or 'Not recorded'}</p>"
            f"<p><strong>Assertion outcome:</strong> {escape(_assertion_outcome(step))}</p>"
            f"<p><strong>Error:</strong> {escape(_text(step.error_details)) or 'None recorded'}</p>"
            f"<p><strong>Screenshot:</strong> {screenshot_html}</p>"
            f"<p><strong>Timestamp:</strong> {escape(step.created_at.isoformat())}</p>"
            "</article>"
        )
    if not rows:
        rows.append("<p>No test steps were recorded.</p>")

    bugs = []
    for bug in data["bugs"]:
        evidence = resolve_evidence_path(run_id, bug.evidence_path)
        evidence_html = "No screenshot evidence recorded."
        if evidence:
            evidence_html = (
                "<span>Evidence file recorded for this bug.</span>"
            )
        bugs.append(
            "<article class=\"step\">"
            f"<h3>{escape(_text(bug.title))}</h3>"
            f"<p><strong>Severity:</strong> {escape(bug.severity)}</p>"
            f"<p>{escape(_text(bug.description))}</p>"
            f"<p><strong>Expected:</strong> {escape(_text(bug.expected_result)) or 'Not recorded'}</p>"
            f"<p><strong>Actual:</strong> {escape(_text(bug.actual_result)) or 'Not recorded'}</p>"
            f"<p><strong>Evidence:</strong> {evidence_html}</p>"
            "</article>"
        )
    if not bugs:
        bugs.append("<p>No bug reports were recorded.</p>")

    logs = []
    for log in data["logs"]:
        logs.append(
            "<li>"
            f"{escape(log.timestamp.isoformat())} — "
            f"{escape(_text(log.event_type))}: {escape(_text(log.message))}"
            "</li>"
        )
    if not logs:
        logs.append("<li>No execution logs were recorded.</li>")

    browser_items = []
    for error in data.get("browser_errors") or []:
        step_label = f"Step {error.step_number}: " if error.step_number is not None else ""
        status_label = f" (Status {error.http_status})" if error.http_status else ""
        type_label = f" [{escape(error.resource_type)}]" if error.resource_type else ""
        browser_items.append(
            "<article class=\"step\">"
            f"<h3>{escape(step_label)}{escape(error.category)}{status_label}{type_label}</h3>"
            f"<p><strong>Message:</strong> {escape(_text(error.message))}</p>"
            f"<p><strong>URL:</strong> {escape(_text(error.url)) or 'None recorded'}</p>"
            f"<p><strong>Timestamp:</strong> {escape(error.timestamp.isoformat())}</p>"
            "</article>"
        )
    if not browser_items:
        browser_items.append("<p>No browser errors were detected.</p>")

    ai_items = []
    for rep in data.get("ai_bug_reports") or []:
        ai_items.append(
            "<article class=\"step\">"
            f"<h3>{escape(_text(rep.title))}</h3>"
            f"<p><strong>Model:</strong> {escape(_text(rep.model))}</p>"
            f"<p><strong>Generated:</strong> {escape(rep.created_at.isoformat())}</p>"
            f"<pre style=\"white-space:pre-wrap;font-family:inherit;\">{escape(_text(rep.report_markdown))}</pre>"
            "</article>"
        )

    failure_summary = _failure_summary(run, data["steps"])
    content = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>QAI Test Report — {escape(_text(run.run_name))}</title>
<style>
body{{font:16px Arial,sans-serif;max-width:1000px;margin:32px auto;padding:0 18px;color:#1f2937}}
.step{{border:1px solid #d1d5db;border-radius:8px;padding:14px;margin:12px 0}}
dt{{font-weight:bold;margin-top:8px}} dd{{margin-left:0}}
</style>
</head>
<body>
<h1>Test Execution Report</h1>
<dl>
<dt>Run ID</dt><dd>{escape(run_id)}</dd>
<dt>Name</dt><dd>{escape(_text(run.run_name))}</dd>
<dt>Platform</dt><dd>{escape(run.platform)}</dd>
<dt>Mode</dt><dd>{escape(run.mode)}</dd>
<dt>Status</dt><dd>{escape(run.status)}</dd>
<dt>Target</dt><dd>{escape(_text(run.target))}</dd>
<dt>Assertion result</dt><dd>{escape(run.assertion_result)}</dd>
<dt>Suspected bug count</dt><dd>{run.suspected_bug_count}</dd>
<dt>Assertions passed</dt><dd>{run.assertion_passed_count}</dd>
<dt>Assertions failed</dt><dd>{run.assertion_failed_count}</dd>
<dt>Execution error</dt><dd>{escape(run.execution_error) if run.execution_error else "None"}</dd>
<dt>Created</dt><dd>{escape(run.created_at.isoformat())}</dd>
<dt>Started</dt><dd>{escape(run.started_at.isoformat()) if run.started_at else "Not started"}</dd>
<dt>Ended</dt><dd>{escape(run.ended_at.isoformat()) if run.ended_at else "Not ended"}</dd>
<dt>Duration</dt><dd>{_duration_ms(run.started_at, run.ended_at)} ms</dd>
<dt>Failure summary</dt><dd>{escape(_text(failure_summary)) or "None recorded"}</dd>
<dt>Goal</dt><dd>{escape(_text(run.goal))}</dd>
<dt>Expected results</dt><dd>{escape(_text("; ".join(run.expected_results))) or "None recorded"}</dd>
</dl>
<h2>Test Steps</h2>
{''.join(rows)}
<h2>Browser Findings</h2>
{''.join(browser_items)}
<h2>AI Failure Analysis</h2>
{''.join(ai_items) if ai_items else '<p>No AI failure analysis recorded.</p>'}
<h2>Bug Reports</h2>
{''.join(bugs)}
<h2>Execution Logs</h2>
<ul>{''.join(logs)}</ul>
</body>
</html>"""
    output_path = _report_directory(run_id) / "report.html"
    output_path.write_text(content, encoding="utf-8")
    return output_path


def _duration_ms(started_at: Any, ended_at: Any) -> int | None:
    if started_at is None or ended_at is None:
        return None
    return max(0, int((ended_at - started_at).total_seconds() * 1000))


def _failure_summary(run: Any, steps: list[Any]) -> str | None:
    if run.error_details:
        return _text(run.error_details)
    for step in steps:
        if step.status in {"failed", "error"}:
            return _text(step.error_details or step.actual_result)
    return None
