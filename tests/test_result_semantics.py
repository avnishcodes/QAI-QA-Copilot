import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from fastapi.testclient import TestClient
from unittest.mock import patch

from backend import database, reports
from backend.app import app
from backend.security import API_TOKEN
from backend.schemas import BugRecordCreate, TestRunCreate, TestStepCreate


def create_run(path: Path, mode: str = "user-directed"):
    database.initialize_database(path)
    return database.create_test_run(
        TestRunCreate(
            run_name="Result semantics",
            platform="website",
            mode=mode,
            target="https://example.test/login",
            goal="Verify the configured assertion.",
            login_username_selector="#username" if mode == "user-directed" else None,
            login_password_selector="#password" if mode == "user-directed" else None,
            login_submit_selector="button[type='submit']" if mode == "user-directed" else None,
            rejection_selector=".error" if mode == "user-directed" else None,
        ),
        path,
    )


class TestResultSemantics(unittest.TestCase):
    def test_completed_run_with_all_assertions_passed(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "results.db"
            run = create_run(path)
            database.update_test_run_status(run.id, "running", path)
            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="check_visibility",
                    status="passed",
                ),
                path,
            )
            database.update_test_run_status(run.id, "completed", path)

            result = database.get_test_run(run.id, path)
            self.assertEqual(result.status, "completed")
            self.assertEqual(result.assertion_result, "passed")
            self.assertEqual(result.suspected_bug_count, 0)
            self.assertIsNone(result.execution_error)

    def test_completed_run_with_failed_assertion_is_not_a_pass(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "results.db"
            run = create_run(path)
            database.update_test_run_status(run.id, "running", path)
            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="check_visibility",
                    status="failed",
                    error_details="Expected rejection was not visible.",
                ),
                path,
            )
            database.update_test_run_status(run.id, "completed", path)

            result = database.get_test_run(run.id, path)
            self.assertEqual(result.status, "completed")
            self.assertEqual(result.assertion_result, "failed")
            self.assertIsNone(result.execution_error)

    def test_execution_error_categories_are_distinguished(self):
        cases = (
            ("failed", "runtime_limit", "Configured execution timeout was reached.", "timeout"),
            ("failed", "browser_error", "Website action failed (TimeoutError).", "browser"),
            ("failed", "safety_blocked", "Unsafe navigation was blocked.", "safety"),
            ("failed", "infrastructure_error", "Website test could not be completed.", "infrastructure"),
        )
        with TemporaryDirectory() as directory:
            path = Path(directory) / "errors.db"
            for index, (status, event_type, message, expected) in enumerate(cases):
                run = create_run(path)
                database.update_test_run_status(run.id, "running", path)
                database.add_execution_log(run.id, event_type, message, database_path=path)
                database.set_test_run_error(run.id, message, path)
                database.update_test_run_status(run.id, status, path)
                result = database.get_test_run(run.id, path)
                self.assertEqual(result.execution_error, expected, index)

    def test_cancelled_and_no_assertion_runs_remain_distinguishable(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "results.db"
            cancelled = create_run(path)
            database.update_test_run_status(cancelled.id, "running", path)
            database.update_test_run_status(cancelled.id, "cancelled", path)
            cancelled_result = database.get_test_run(cancelled.id, path)
            self.assertEqual(cancelled_result.status, "cancelled")
            self.assertEqual(cancelled_result.assertion_result, "not_evaluated")
            self.assertIsNone(cancelled_result.execution_error)

            empty = create_run(path)
            database.update_test_run_status(empty.id, "running", path)
            database.update_test_run_status(empty.id, "completed", path)
            empty_result = database.get_test_run(empty.id, path)
            self.assertEqual(empty_result.assertion_result, "not_evaluated")

    def test_autonomous_failed_assertion_counts_only_persisted_suspected_bugs(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "results.db"
            run = create_run(path, "autonomous-exploration")
            database.update_test_run_status(run.id, "running", path)
            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="autonomous:assert_visible",
                    status="failed",
                    error_details="Assertion failed.",
                ),
                path,
            )
            database.add_bug_record(
                run.id,
                BugRecordCreate(
                    title="Suspected missing element",
                    description="The observed element was not visible.",
                    severity="Medium",
                    confidence=0.5,
                ),
                path,
            )
            database.update_test_run_status(run.id, "completed", path)

            result = database.get_test_run(run.id, path)
            self.assertEqual(result.assertion_result, "failed")
            self.assertEqual(result.suspected_bug_count, 1)

    def test_history_and_reports_include_additive_result_fields(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "history.db"
            run = create_run(path)
            database.update_test_run_status(run.id, "running", path)
            database.add_test_step(
                run.id,
                TestStepCreate(step_number=1, action="check_visibility", status="failed"),
                path,
            )
            database.update_test_run_status(run.id, "completed", path)

            history, total = database.list_test_run_history(database_path=path)
            self.assertEqual(total, 1)
            self.assertEqual(history[0].assertion_result, "failed")
            self.assertEqual(history[0].suspected_bug_count, 0)

            data = database.get_test_run_report_data(run.id, path)
            old_reports = reports.REPORTS_DIR
            old_project = reports.PROJECT_DIR
            try:
                reports.REPORTS_DIR = root / "reports"
                reports.PROJECT_DIR = root
                html = reports.generate_html_report(data, "http://127.0.0.1:8000")
                workbook = load_workbook(
                    reports.generate_excel_report(data), read_only=True
                )
                summary = "\n".join(
                    str(row[0]) + ":" + str(row[1])
                    for row in workbook["Test Summary"].iter_rows(values_only=True)
                )
                workbook.close()
            finally:
                reports.REPORTS_DIR = old_reports
                reports.PROJECT_DIR = old_project

            html_text = html.read_text(encoding="utf-8")
            for text in ("Assertion result", "failed", "Suspected bug count"):
                self.assertIn(text, html_text + summary)

    def test_api_run_and_history_responses_expose_result_semantics(self):
        with TemporaryDirectory() as directory:
            old_path = database.DATABASE_PATH
            database.DATABASE_PATH = Path(directory) / "api.db"
            try:
                with TestClient(app, headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"}) as client:
                    created = client.post(
                        "/test-runs",
                        json={
                            "run_name": "API result semantics",
                            "platform": "website",
                            "mode": "user-directed",
                            "target": "https://example.test/login",
                            "goal": "Verify the assertion result.",
                        },
                    )
                    self.assertEqual(created.status_code, 201)
                    run_id = created.json()["id"]
                    self.assertEqual(
                        created.json()["assertion_result"], "not_evaluated"
                    )
                    self.assertEqual(created.json()["suspected_bug_count"], 0)

                    with patch(
                        "backend.app.validate_target_url",
                        return_value="example.test",
                    ), patch("backend.app.build_login_action_plan"):
                        run = client.get(f"/test-runs/{run_id}")
                        history = client.get("/test-runs?limit=1")
                    self.assertEqual(run.status_code, 200)
                    self.assertIn("execution_error", run.json())
                    self.assertEqual(history.status_code, 200)
                    self.assertIn(
                        "assertion_result", history.json()["items"][0]
                    )
            finally:
                database.DATABASE_PATH = old_path


if __name__ == "__main__":
    unittest.main()
