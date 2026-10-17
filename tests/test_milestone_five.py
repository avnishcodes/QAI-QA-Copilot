import io
import json
import os
import socket
import subprocess
import sys
import threading
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

os.environ.setdefault("GEMINI_API_KEY", "milestone-test-key")

from fastapi.testclient import TestClient

from backend import database, reports, website_runner
from backend import app as app_module
from backend.security import API_TOKEN


FRONTEND_URL = (Path(__file__).resolve().parent.parent / "frontend" / "index.html").as_uri()


class TestFrontendReportEscaping(unittest.TestCase):
    def test_gemini_report_and_error_text_cannot_inject_html(self):
        malicious_text = '<img src=x onerror="window.__qaiXss = 1">'
        responses = [
            {"success": True, "report": malicious_text},
            {"success": False, "error": malicious_text},
        ]

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            page.add_init_script("window.__qaiXss = 0")
            current_response = {"value": responses.pop(0)}

            def respond(route):
                route.fulfill(
                    status=200,
                    content_type="application/json",
                    body=json.dumps(current_response["value"]),
                )

            page.route("http://127.0.0.1:8000/analyze-bug", respond)
            page.goto(FRONTEND_URL)

            test_responses = [current_response["value"], *responses]
            for expected_response in test_responses:
                current_response["value"] = expected_response
                page.locator("#bugInput").fill("A button does not respond.")
                page.locator("#analyzeBtn").click()
                page.wait_for_function(
                    "text => document.querySelector('#result').textContent.includes(text)",
                    arg=malicious_text,
                )

                self.assertEqual(page.locator("#result img").count(), 0)
                self.assertEqual(page.evaluate("window.__qaiXss"), 0)
                self.assertIn(malicious_text, page.locator("#result").inner_text())

            browser.close()


class TestGeminiFailureHandling(unittest.TestCase):
    def test_backend_module_can_start_without_a_gemini_key(self):
        environment = os.environ.copy()
        environment.pop("GEMINI_API_KEY", None)
        environment["PYTHON_DOTENV_DISABLED"] = "1"
        result = subprocess.run(
            [
                sys.executable,
                "-c",
                "from backend.app import app; print(app.title)",
            ],
            cwd=Path(__file__).resolve().parent.parent,
            env=environment,
            capture_output=True,
            text=True,
            timeout=15,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), "QAI - QA Copilot")

    def test_missing_key_preserves_api_and_reports_configuration_error(self):
        with TemporaryDirectory() as temporary_directory:
            original_database_path = database.DATABASE_PATH
            database.DATABASE_PATH = Path(temporary_directory) / "no-key.db"
            try:
                with patch.object(app_module, "client", None):
                    with TestClient(app_module.app, headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"}) as client:
                        self.assertEqual(client.get("/").status_code, 200)
                        response = client.post(
                            "/analyze-bug",
                            json={"observation": "The save button does not respond."},
                        )
                    with self.assertRaisesRegex(RuntimeError, "not configured"):
                        app_module.generate_autonomous_plan("unused")
            finally:
                database.DATABASE_PATH = original_database_path

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["success"])
        self.assertIn("GEMINI_API_KEY is not configured", response.json()["error"])

    def test_provider_exception_details_are_not_returned_or_logged(self):
        provider_detail = "503 UNAVAILABLE api_key=private-marker"
        with TemporaryDirectory() as temporary_directory:
            original_database_path = database.DATABASE_PATH
            database.DATABASE_PATH = Path(temporary_directory) / "gemini-error.db"
            output = io.StringIO()
            try:
                with TestClient(app_module.app, headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"}) as client:
                    with (
                        patch.object(
                            app_module.client.models,
                            "generate_content",
                            side_effect=RuntimeError(provider_detail),
                        ) as generate_content,
                        patch("backend.app.time.sleep"),
                        redirect_stdout(output),
                    ):
                        response = client.post(
                            "/analyze-bug",
                            json={"observation": "The save button does not respond."},
                        )
                self.assertEqual(generate_content.call_count, 3)
            finally:
                database.DATABASE_PATH = original_database_path

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["success"])
        self.assertNotIn("private-marker", response.text)
        self.assertNotIn("private-marker", output.getvalue())


LOGIN_FIXTURE_HTML = """<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>Secure Login Fixture</title></head>
<body>
  <h1>Authentication Portal</h1>
  <form id="login-form">
    <label for="username">Username</label>
    <input id="username" name="username" type="text" autocomplete="username">
    <label for="password">Password</label>
    <input id="password" name="password" type="password" autocomplete="current-password">
    <button id="submit-btn" type="submit">Sign in</button>
  </form>
  <div class="error-message" hidden></div>
  <div class="account-menu" hidden>User Account Dashboard</div>
  <script>
    document.querySelector('#login-form').addEventListener('submit', function(e) {
      e.preventDefault();
      var err = document.querySelector('.error-message');
      err.textContent = 'Invalid credentials';
      err.hidden = false;
    });
  </script>
</body>
</html>"""

AUTONOMOUS_FIXTURE_HTML = """<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>Portal Explorer</title></head>
<body>
  <h1>Portal Main</h1>
  <input id="search-box" type="search" placeholder="Search documentation" aria-label="Search catalogue">
  <form id="filter-form" action="/search" method="get">
    <input id="filter-query" type="text" name="query" placeholder="Filter query">
    <button type="submit">Apply filter</button>
  </form>
  <a href="/approved-docs">Approved Docs</a>
  <a href="/delete-account">Delete Account</a>
  <a href="/logout">Logout</a>
  <form id="danger-form" action="/destroy" method="post">
    <input type="password" name="secret_token">
    <button type="submit">Destroy Data</button>
  </form>
</body>
</html>"""

SEARCH_FIXTURE_HTML = """<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>Search Results</title></head>
<body>
  <h1>Search Results</h1>
  <p>Search completed successfully.</p>
</body>
</html>"""

DOCS_FIXTURE_HTML = """<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>Approved Documentation</title></head>
<body>
  <h1>Documentation Portal</h1>
  <p>Official system documentation.</p>
</body>
</html>"""


def make_e2e_route_interceptor(fixture_html=LOGIN_FIXTURE_HTML, slow_event=None):
    def interceptor(route):
        url = route.request.url
        if "fixture.qa.example.com" not in url:
            route.abort("blockedbyclient")
            return True
        if slow_event is not None:
            slow_event.set()
            time.sleep(0.4)
        if "/search" in url:
            route.fulfill(status=200, content_type="text/html", body=SEARCH_FIXTURE_HTML)
            return True
        if "/approved-docs" in url:
            route.fulfill(status=200, content_type="text/html", body=DOCS_FIXTURE_HTML)
            return True
        if "/login" in url or "/explore" in url or url.endswith(".example.com/"):
            body = fixture_html
            route.fulfill(status=200, content_type="text/html", body=body)
            return True
        route.abort("blockedbyclient")
        return True
    return interceptor


class TestEndToEndBrowserValidation(unittest.TestCase):
    def test_e2e_user_directed_login_complete_lifecycle(self):
        with TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            original_evidence_dir = website_runner.EVIDENCE_DIR
            original_runner_project_dir = website_runner.PROJECT_DIR
            original_reports_dir = reports.REPORTS_DIR
            original_reports_project_dir = reports.PROJECT_DIR

            database.DATABASE_PATH = temporary_root / "e2e_login.db"
            website_runner.PROJECT_DIR = temporary_root
            website_runner.EVIDENCE_DIR = temporary_root / "evidence"
            reports.PROJECT_DIR = temporary_root
            reports.REPORTS_DIR = temporary_root / "generated_reports"

            interceptor = make_e2e_route_interceptor(LOGIN_FIXTURE_HTML)

            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[
                            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))
                        ],
                    ),
                    patch.dict(
                        website_runner.execute_website_run.__kwdefaults__,
                        {"request_interceptor": interceptor},
                    ),
                ):
                    with TestClient(
                        app_module.app,
                        headers={
                            "Host": "127.0.0.1:8000",
                            "Authorization": f"Bearer {API_TOKEN}",
                        },
                    ) as client:
                        # 1. Create run via API
                        create_response = client.post(
                            "/test-runs",
                            json={
                                "run_name": "E2E Login Verification",
                                "platform": "website",
                                "mode": "user-directed",
                                "target": "https://fixture.qa.example.com/login",
                                "goal": "Verify invalid login credentials trigger rejection.",
                                "expected_results": ["Rejection message is visible."],
                                "max_steps": 10,
                                "login_username_selector": "#username",
                                "login_password_selector": "#password",
                                "login_submit_selector": "#submit-btn",
                                "rejection_selector": ".error-message",
                                "rejection_text": "Invalid credentials",
                                "authenticated_selector": ".account-menu",
                            },
                        )
                        self.assertEqual(create_response.status_code, 201)
                        run_id = create_response.json()["id"]

                        # 2. Start run via API (triggers real Chromium in background task)
                        start_response = client.post(
                            f"/test-runs/{run_id}/start",
                            json={
                                "username": "test-user@example.invalid",
                                "password": "synthetic-test-password",
                            },
                        )
                        self.assertEqual(start_response.status_code, 202)

                        # 3. Verify run status and assertion outcome in SQLite
                        run_data = client.get(f"/test-runs/{run_id}").json()
                        self.assertEqual(run_data["status"], "completed")
                        self.assertEqual(run_data["assertion_result"], "passed")
                        self.assertIsNone(run_data["execution_error"])

                        # 4. Verify step records and timestamps
                        steps = client.get(f"/test-runs/{run_id}/steps").json()
                        self.assertEqual(len(steps), 7)
                        self.assertTrue(all(step["status"] == "passed" for step in steps))
                        self.assertTrue(all(step["duration_ms"] is not None for step in steps))

                        # 5. Verify dashboard metrics
                        metrics = client.get("/dashboard/metrics").json()
                        self.assertGreaterEqual(metrics["total_runs"], 1)
                        self.assertGreaterEqual(metrics["completed_runs"], 1)
                        self.assertGreaterEqual(metrics["assertions_passed"], 1)

                        # 6. Verify HTML report generation and credential redaction
                        html_response = client.get(f"/test-runs/{run_id}/reports/html")
                        self.assertEqual(html_response.status_code, 200)
                        self.assertIn("text/html", html_response.headers["content-type"])
                        self.assertIn("E2E Login Verification", html_response.text)
                        self.assertNotIn("synthetic-test-password", html_response.text)

                        # 7. Verify Excel report generation and schema
                        excel_response = client.get(f"/test-runs/{run_id}/reports/excel")
                        self.assertEqual(excel_response.status_code, 200)
                        report_workbook = load_workbook(io.BytesIO(excel_response.content))
                        self.assertIn("Test Summary", report_workbook.sheetnames)
                        self.assertIn("Test Steps", report_workbook.sheetnames)
                        self.assertIn("Bug Reports", report_workbook.sheetnames)
                        self.assertIn("Execution Logs", report_workbook.sheetnames)

                        # 8. Dynamically locate screenshot step and download evidence
                        screenshot_step = next(s for s in steps if s.get("screenshot_path"))
                        screenshot_response = client.get(
                            f"/test-runs/{run_id}/steps/{screenshot_step['step_number']}/screenshot"
                        )
                        self.assertEqual(screenshot_response.status_code, 200)
                        self.assertEqual(screenshot_response.headers["content-type"], "image/png")
                        self.assertTrue(screenshot_response.content.startswith(b"\x89PNG\r\n\x1a\n"))
            finally:
                database.DATABASE_PATH = original_database_path
                website_runner.EVIDENCE_DIR = original_evidence_dir
                website_runner.PROJECT_DIR = original_runner_project_dir
                reports.REPORTS_DIR = original_reports_dir
                reports.PROJECT_DIR = original_reports_project_dir

    def test_e2e_autonomous_exploration_complete_lifecycle(self):
        with TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            original_evidence_dir = website_runner.EVIDENCE_DIR
            original_runner_project_dir = website_runner.PROJECT_DIR
            original_reports_dir = reports.REPORTS_DIR
            original_reports_project_dir = reports.PROJECT_DIR

            database.DATABASE_PATH = temporary_root / "e2e_auto.db"
            website_runner.PROJECT_DIR = temporary_root
            website_runner.EVIDENCE_DIR = temporary_root / "evidence"
            reports.PROJECT_DIR = temporary_root
            reports.REPORTS_DIR = temporary_root / "generated_reports"

            interceptor = make_e2e_route_interceptor(AUTONOMOUS_FIXTURE_HTML)

            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[
                            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))
                        ],
                    ),
                    patch.dict(
                        website_runner.execute_autonomous_website_run.__kwdefaults__,
                        {"request_interceptor": interceptor},
                    ),
                    patch(
                        "backend.app.generate_autonomous_plan",
                        return_value=None,
                    ),
                ):
                    with TestClient(
                        app_module.app,
                        headers={
                            "Host": "127.0.0.1:8000",
                            "Authorization": f"Bearer {API_TOKEN}",
                        },
                    ) as client:
                        create_response = client.post(
                            "/test-runs",
                            json={
                                "run_name": "E2E Autonomous Exploration",
                                "platform": "website",
                                "mode": "autonomous-exploration",
                                "target": "https://fixture.qa.example.com/explore",
                                "goal": "Safely inspect portal and search field.",
                                "max_steps": 6,
                                "max_duration_seconds": 30,
                                "allow_form_submission": False,
                            },
                        )
                        self.assertEqual(create_response.status_code, 201)
                        run_id = create_response.json()["id"]

                        start_response = client.post(f"/test-runs/{run_id}/start")
                        self.assertEqual(start_response.status_code, 202)

                        run_data = client.get(f"/test-runs/{run_id}").json()
                        self.assertEqual(run_data["status"], "completed")

                        steps = client.get(f"/test-runs/{run_id}/steps").json()
                        self.assertGreaterEqual(len(steps), 2)
                        self.assertEqual(steps[0]["action"], "autonomous:inspect")

                        # Verify metrics updated
                        metrics = client.get("/dashboard/metrics").json()
                        self.assertGreaterEqual(metrics["completed_runs"], 1)

                        # Verify HTML report download
                        html_response = client.get(f"/test-runs/{run_id}/reports/html")
                        self.assertEqual(html_response.status_code, 200)
            finally:
                database.DATABASE_PATH = original_database_path
                website_runner.EVIDENCE_DIR = original_evidence_dir
                website_runner.PROJECT_DIR = original_runner_project_dir
                reports.REPORTS_DIR = original_reports_dir
                reports.PROJECT_DIR = original_reports_project_dir

    def test_e2e_assertion_failure_registers_failed_outcome(self):
        with TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            original_evidence_dir = website_runner.EVIDENCE_DIR
            original_runner_project_dir = website_runner.PROJECT_DIR
            original_reports_dir = reports.REPORTS_DIR
            original_reports_project_dir = reports.PROJECT_DIR

            database.DATABASE_PATH = temporary_root / "e2e_fail.db"
            website_runner.PROJECT_DIR = temporary_root
            website_runner.EVIDENCE_DIR = temporary_root / "evidence"
            reports.PROJECT_DIR = temporary_root
            reports.REPORTS_DIR = temporary_root / "generated_reports"

            interceptor = make_e2e_route_interceptor(LOGIN_FIXTURE_HTML)

            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[
                            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))
                        ],
                    ),
                    patch.dict(
                        website_runner.execute_website_run.__kwdefaults__,
                        {"request_interceptor": interceptor},
                    ),
                ):
                    with TestClient(
                        app_module.app,
                        headers={
                            "Host": "127.0.0.1:8000",
                            "Authorization": f"Bearer {API_TOKEN}",
                        },
                    ) as client:
                        create_response = client.post(
                            "/test-runs",
                            json={
                                "run_name": "E2E Failed Assertion Test",
                                "platform": "website",
                                "mode": "user-directed",
                                "target": "https://fixture.qa.example.com/login",
                                "goal": "Verify assertion failure recording.",
                                "max_steps": 10,
                                "login_username_selector": "#username",
                                "login_password_selector": "#password",
                                "login_submit_selector": "#submit-btn",
                                "rejection_selector": ".error-message",
                                "rejection_text": "Nonexistent expected message",
                            },
                        )
                        self.assertEqual(create_response.status_code, 201)
                        run_id = create_response.json()["id"]

                        start_response = client.post(f"/test-runs/{run_id}/start")
                        self.assertEqual(start_response.status_code, 202)

                        run_data = client.get(f"/test-runs/{run_id}").json()
                        self.assertEqual(run_data["status"], "completed")
                        self.assertEqual(run_data["assertion_result"], "failed")
                        self.assertEqual(run_data["assertion_failed_count"], 1)

                        # Report reflects failed assertion
                        html_response = client.get(f"/test-runs/{run_id}/reports/html")
                        self.assertEqual(html_response.status_code, 200)
                        self.assertIn("failed", html_response.text.lower())
            finally:
                database.DATABASE_PATH = original_database_path
                website_runner.EVIDENCE_DIR = original_evidence_dir
                website_runner.PROJECT_DIR = original_runner_project_dir
                reports.REPORTS_DIR = original_reports_dir
                reports.PROJECT_DIR = original_reports_project_dir

    def test_e2e_cancellation_halts_active_browser_run(self):
        with TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            original_evidence_dir = website_runner.EVIDENCE_DIR
            original_runner_project_dir = website_runner.PROJECT_DIR
            original_reports_dir = reports.REPORTS_DIR
            original_reports_project_dir = reports.PROJECT_DIR

            database.DATABASE_PATH = temporary_root / "e2e_cancel.db"
            website_runner.PROJECT_DIR = temporary_root
            website_runner.EVIDENCE_DIR = temporary_root / "evidence"
            reports.PROJECT_DIR = temporary_root
            reports.REPORTS_DIR = temporary_root / "generated_reports"

            slow_event = threading.Event()
            interceptor = make_e2e_route_interceptor(LOGIN_FIXTURE_HTML, slow_event=slow_event)

            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[
                            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))
                        ],
                    ),
                    patch.dict(
                        website_runner.execute_website_run.__kwdefaults__,
                        {"request_interceptor": interceptor},
                    ),
                ):
                    with TestClient(
                        app_module.app,
                        headers={
                            "Host": "127.0.0.1:8000",
                            "Authorization": f"Bearer {API_TOKEN}",
                        },
                    ) as client:
                        create_response = client.post(
                            "/test-runs",
                            json={
                                "run_name": "E2E Cancel Test",
                                "platform": "website",
                                "mode": "user-directed",
                                "target": "https://fixture.qa.example.com/login",
                                "goal": "Verify run cancellation.",
                                "max_steps": 10,
                                "login_username_selector": "#username",
                                "login_password_selector": "#password",
                                "login_submit_selector": "#submit-btn",
                                "rejection_selector": ".error-message",
                            },
                        )
                        run_id = create_response.json()["id"]

                        def trigger_cancel():
                            slow_event.wait(timeout=5)
                            time.sleep(0.05)
                            cancel_response = client.post(f"/test-runs/{run_id}/cancel")
                            self.assertEqual(cancel_response.status_code, 200)

                        cancel_thread = threading.Thread(target=trigger_cancel)
                        cancel_thread.start()

                        start_response = client.post(f"/test-runs/{run_id}/start")
                        self.assertEqual(start_response.status_code, 202)
                        cancel_thread.join(timeout=10)

                        run_data = client.get(f"/test-runs/{run_id}").json()
                        self.assertEqual(run_data["status"], "cancelled")
            finally:
                database.DATABASE_PATH = original_database_path
                website_runner.EVIDENCE_DIR = original_evidence_dir
                website_runner.PROJECT_DIR = original_runner_project_dir
                reports.REPORTS_DIR = original_reports_dir
                reports.PROJECT_DIR = original_reports_project_dir


class TestEndToEndNegativeSecurity(unittest.TestCase):
    def test_unauthorized_api_requests_are_rejected(self):
        with TestClient(app_module.app, headers={"Host": "127.0.0.1:8000"}) as client:
            # Missing token
            self.assertEqual(client.get("/test-runs").status_code, 401)
            self.assertEqual(
                client.post(
                    "/test-runs",
                    json={
                        "run_name": "unauthorized",
                        "platform": "website",
                        "mode": "user-directed",
                        "target": "https://example.test",
                        "goal": "unauthorized",
                    },
                ).status_code,
                401,
            )
            # Invalid token
            self.assertEqual(
                client.get("/test-runs", headers={"Authorization": "Bearer bad-token"}).status_code,
                401,
            )
            # Invalid origin on state-changing request
            self.assertEqual(
                client.post(
                    "/test-runs",
                    json={
                        "run_name": "bad origin",
                        "platform": "website",
                        "mode": "user-directed",
                        "target": "https://example.test",
                        "goal": "bad origin",
                    },
                    headers={"Authorization": f"Bearer {API_TOKEN}", "Origin": "https://evil.attacker.com"},
                ).status_code,
                403,
            )
            # Invalid host
            self.assertEqual(
                client.get(
                    "/test-runs",
                    headers={"Host": "malicious.host:8000", "Authorization": f"Bearer {API_TOKEN}"},
                ).status_code,
                400,
            )

    def test_ssrf_and_private_network_targets_are_rejected(self):
        targets = [
            "http://127.0.0.1/admin",
            "http://localhost:8000/",
            "http://169.254.169.254/latest/meta-data/",
            "http://168.63.129.16/",
            "ftp://fixture.qa.example.com/",
            "https://fixture.qa.example.com:8443/",
        ]
        with TestClient(
            app_module.app,
            headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"},
        ) as client:
            for target in targets:
                with self.subTest(target=target):
                    response = client.post(
                        "/test-runs",
                        json={
                            "run_name": "SSRF Test",
                            "platform": "website",
                            "mode": "user-directed",
                            "target": target,
                            "goal": "Verify SSRF block.",
                        },
                    )
                    self.assertEqual(response.status_code, 422)

    def test_hostname_resolving_to_private_address_is_rejected_at_start(self):
        with TemporaryDirectory() as temporary_directory:
            original_database_path = database.DATABASE_PATH
            database.DATABASE_PATH = Path(temporary_directory) / "dns_block.db"
            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("10.0.0.1", 443))],
                    ),
                    TestClient(
                        app_module.app,
                        headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"},
                    ) as client,
                ):
                    create_response = client.post(
                        "/test-runs",
                        json={
                            "run_name": "Private IP Test",
                            "platform": "website",
                            "mode": "user-directed",
                            "target": "https://internal-spoof.example.com/login",
                            "goal": "Verify DNS resolution block.",
                            "login_username_selector": "#user",
                            "login_password_selector": "#pass",
                            "login_submit_selector": "#btn",
                            "rejection_selector": ".err",
                        },
                    )
                    run_id = create_response.json()["id"]
                    start_response = client.post(f"/test-runs/{run_id}/start")
                    self.assertEqual(start_response.status_code, 422)
            finally:
                database.DATABASE_PATH = original_database_path

    def test_autonomous_mode_blocks_destructive_paths_and_sensitive_fields(self):
        with TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            original_evidence_dir = website_runner.EVIDENCE_DIR
            original_runner_project_dir = website_runner.PROJECT_DIR
            original_reports_dir = reports.REPORTS_DIR
            original_reports_project_dir = reports.PROJECT_DIR

            database.DATABASE_PATH = temporary_root / "e2e_destructive.db"
            website_runner.PROJECT_DIR = temporary_root
            website_runner.EVIDENCE_DIR = temporary_root / "evidence"
            reports.PROJECT_DIR = temporary_root
            reports.REPORTS_DIR = temporary_root / "generated_reports"

            interceptor = make_e2e_route_interceptor(AUTONOMOUS_FIXTURE_HTML)

            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))],
                    ),
                    patch.dict(
                        website_runner.execute_autonomous_website_run.__kwdefaults__,
                        {"request_interceptor": interceptor},
                    ),
                    patch(
                        "backend.app.generate_autonomous_plan",
                        return_value=None,
                    ),
                ):
                    with TestClient(
                        app_module.app,
                        headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"},
                    ) as client:
                        create_response = client.post(
                            "/test-runs",
                            json={
                                "run_name": "Destructive Action Block Test",
                                "platform": "website",
                                "mode": "autonomous-exploration",
                                "target": "https://fixture.qa.example.com/explore",
                                "goal": "Verify destructive actions are ignored.",
                                "max_steps": 10,
                                "max_duration_seconds": 30,
                                "allow_form_submission": True,
                            },
                        )
                        run_id = create_response.json()["id"]
                        start_response = client.post(f"/test-runs/{run_id}/start")
                        self.assertEqual(start_response.status_code, 202)

                        steps = client.get(f"/test-runs/{run_id}/steps").json()
                        step_actions = [s.get("page_or_screen", "") + " " + s.get("actual_result", "") for s in steps]
                        for recorded in step_actions:
                            self.assertNotIn("/delete-account", recorded)
                            self.assertNotIn("/logout", recorded)
                            self.assertNotIn("/destroy", recorded)
                            self.assertNotIn("secret_token", recorded)
            finally:
                database.DATABASE_PATH = original_database_path
                website_runner.EVIDENCE_DIR = original_evidence_dir
                website_runner.PROJECT_DIR = original_runner_project_dir
                reports.REPORTS_DIR = original_reports_dir
                reports.PROJECT_DIR = original_reports_project_dir

    def test_autonomous_form_submission_denied_without_explicit_opt_in(self):
        with TemporaryDirectory() as temporary_directory:
            temporary_root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            original_evidence_dir = website_runner.EVIDENCE_DIR
            original_runner_project_dir = website_runner.PROJECT_DIR
            original_reports_dir = reports.REPORTS_DIR
            original_reports_project_dir = reports.PROJECT_DIR

            database.DATABASE_PATH = temporary_root / "e2e_form_optin.db"
            website_runner.PROJECT_DIR = temporary_root
            website_runner.EVIDENCE_DIR = temporary_root / "evidence"
            reports.PROJECT_DIR = temporary_root
            reports.REPORTS_DIR = temporary_root / "generated_reports"

            interceptor = make_e2e_route_interceptor(AUTONOMOUS_FIXTURE_HTML)

            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    patch(
                        "backend.website_runner.socket.getaddrinfo",
                        return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))],
                    ),
                    patch.dict(
                        website_runner.execute_autonomous_website_run.__kwdefaults__,
                        {"request_interceptor": interceptor},
                    ),
                    patch(
                        "backend.app.generate_autonomous_plan",
                        return_value=None,
                    ),
                ):
                    with TestClient(
                        app_module.app,
                        headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"},
                    ) as client:
                        # 1. Without opt-in (default allow_form_submission=False)
                        create_response = client.post(
                            "/test-runs",
                            json={
                                "run_name": "Form Denied Without Opt-in",
                                "platform": "website",
                                "mode": "autonomous-exploration",
                                "target": "https://fixture.qa.example.com/explore",
                                "goal": "Verify form submission is not attempted without opt-in.",
                                "max_steps": 5,
                                "max_duration_seconds": 30,
                                "allow_form_submission": False,
                            },
                        )
                        run_id = create_response.json()["id"]
                        client.post(f"/test-runs/{run_id}/start")
                        steps = client.get(f"/test-runs/{run_id}/steps").json()
                        self.assertFalse(any("submit" in s.get("action", "").lower() for s in steps))

                        # 2. With opt-in (allow_form_submission=True)
                        create_opt_in = client.post(
                            "/test-runs",
                            json={
                                "run_name": "Form Allowed With Opt-in",
                                "platform": "website",
                                "mode": "autonomous-exploration",
                                "target": "https://fixture.qa.example.com/explore",
                                "goal": "Verify safe form submission occurs with opt-in.",
                                "max_steps": 6,
                                "max_duration_seconds": 30,
                                "allow_form_submission": True,
                            },
                        )
                        opt_in_id = create_opt_in.json()["id"]
                        client.post(f"/test-runs/{opt_in_id}/start")
                        steps_opt_in = client.get(f"/test-runs/{opt_in_id}/steps").json()
                        self.assertTrue(any("submit_form" in s.get("action", "") for s in steps_opt_in))
            finally:
                database.DATABASE_PATH = original_database_path
                website_runner.EVIDENCE_DIR = original_evidence_dir
                website_runner.PROJECT_DIR = original_runner_project_dir
                reports.REPORTS_DIR = original_reports_dir
                reports.PROJECT_DIR = original_reports_project_dir
