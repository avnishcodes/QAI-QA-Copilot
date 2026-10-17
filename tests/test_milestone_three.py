import os
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from uuid import uuid4

os.environ.setdefault("GEMINI_API_KEY", "milestone-test-key")

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import database, reports
from backend.app import app
from backend.security import API_TOKEN
from backend.schemas import BugRecordCreate, TestRunCreate, TestStepCreate


def create_run(database_path: Path, name: str = "Report test"):
    return database.create_test_run(
        TestRunCreate(
            run_name=name,
            platform="website",
            mode="user-directed",
            target="https://example.test/login",
            goal="Verify invalid login is rejected.",
            expected_results=["Rejection indicator is visible."],
        ),
        database_path,
    )


class TestExecutionHistory(unittest.TestCase):
    def test_history_filters_and_paginates_saved_runs(self):
        with TemporaryDirectory() as temporary_directory:
            database_path = Path(temporary_directory) / "history.db"
            database.initialize_database(database_path)
            website = create_run(database_path, "Website run")
            database.update_test_run_status(website.id, "running", database_path)
            database.update_test_run_status(website.id, "completed", database_path)
            android = database.create_test_run(
                TestRunCreate(
                    run_name="Android run",
                    platform="android",
                    mode="user-directed",
                    target="https://example.test/app",
                    goal="Check app.",
                ),
                database_path,
            )

            first_page, total = database.list_test_run_history(
                limit=1, offset=0, database_path=database_path
            )
            second_page, _ = database.list_test_run_history(
                limit=1, offset=1, database_path=database_path
            )
            filtered, filtered_total = database.list_test_run_history(
                status="completed",
                platform="website",
                database_path=database_path,
            )

            self.assertEqual(total, 2)
            self.assertEqual(len(first_page), 1)
            self.assertEqual(len(second_page), 1)
            self.assertEqual(str(first_page[0].id), android.id)
            self.assertEqual(str(second_page[0].id), website.id)
            self.assertEqual(filtered_total, 1)
            self.assertEqual(filtered[0].status, "completed")
            self.assertIsNotNone(filtered[0].duration_ms)


class TestReportGeneration(unittest.TestCase):
    def test_reports_cover_success_failure_cancelled_and_empty_runs(self):
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            database_path = root / "reports.db"
            database.initialize_database(database_path)
            outputs = root / "reports"
            evidence_root = root / "project"
            evidence_root.mkdir()
            statuses = ("completed", "failed", "cancelled", "pending")
            report_data = {}

            for status in statuses:
                run = create_run(database_path, f"{status} run")
                if status != "pending":
                    database.update_test_run_status(run.id, "running", database_path)
                    if status == "completed":
                        database.add_test_step(
                            run.id,
                            TestStepCreate(
                                step_number=1,
                                action="Verify rejection",
                                actual_result="Rejection was visible.",
                                status="passed",
                            ),
                            database_path,
                        )
                    elif status == "failed":
                        database.add_test_step(
                            run.id,
                            TestStepCreate(
                                step_number=1,
                                action="Verify rejection",
                                actual_result="Rejection was not visible.",
                                error_details="Configured assertion failed.",
                                status="failed",
                            ),
                            database_path,
                        )
                    database.update_test_run_status(run.id, status, database_path)

                report_data[status] = database.get_test_run_report_data(
                    run.id, database_path
                )

            with (
                patch.object(reports, "REPORTS_DIR", outputs),
                patch.object(reports, "PROJECT_DIR", evidence_root),
            ):
                for status, data in report_data.items():
                    excel_path = reports.generate_excel_report(data)
                    html_path = reports.generate_html_report(
                        data, "http://127.0.0.1:8000"
                    )
                    workbook = load_workbook(excel_path, read_only=True)
                    self.assertEqual(
                        workbook.sheetnames,
                        ["Test Summary", "Test Steps", "Bug Reports", "Execution Logs"],
                    )
                    workbook.close()
                    self.assertTrue(html_path.read_text(encoding="utf-8").startswith("<!doctype html>"))
                    self.assertEqual(excel_path.parent.parent, outputs)

                empty_book = load_workbook(
                    reports.generate_excel_report(report_data["pending"]),
                    read_only=True,
                )
                self.assertEqual(empty_book["Test Steps"].max_row, 1)
                self.assertEqual(empty_book["Bug Reports"].max_row, 1)
                empty_book.close()

    def test_html_escaping_evidence_confinement_and_secret_redaction(self):
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            database_path = root / "security.db"
            database.initialize_database(database_path)
            run = create_run(
                database_path,
                '<script>alert("report")</script>',
            )
            database.update_test_run_status(run.id, "running", database_path)

            run_evidence = root / "project" / "evidence" / run.id
            run_evidence.mkdir(parents=True)
            (run_evidence / "step-1.png").write_bytes(b"image fixture")
            foreign_id = str(uuid4())
            foreign_evidence = root / "project" / "evidence" / foreign_id
            foreign_evidence.mkdir(parents=True)
            (foreign_evidence / "foreign.png").write_bytes(b"foreign fixture")

            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="Check <img src=x onerror=alert(1)>",
                    actual_result="password=supersecret token:abc123 Bearer xyz.token",
                    status="failed",
                    error_details="authorization:privatetoken",
                    screenshot_path=f"evidence/{run.id}/step-1.png",
                ),
                database_path,
            )
            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=2,
                    action="Foreign evidence must be omitted",
                    status="passed",
                    screenshot_path=f"evidence/{foreign_id}/foreign.png",
                ),
                database_path,
            )
            database.add_bug_record(
                run.id,
                BugRecordCreate(
                    title="Bug report <script>",
                    description="Potential XSS <img src=x onerror=alert(1)>",
                    severity="High",
                    confidence=0.9,
                ),
                database_path,
            )
            data = database.get_test_run_report_data(run.id, database_path)

            with (
                patch.object(reports, "REPORTS_DIR", root / "reports"),
                patch.object(reports, "PROJECT_DIR", root / "project"),
            ):
                html = reports.generate_html_report(
                    data, "http://127.0.0.1:8000"
                ).read_text(encoding="utf-8")
                excel_path = reports.generate_excel_report(data)
                workbook = load_workbook(excel_path, read_only=True)
                excel_text = "\n".join(
                    str(cell.value)
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None
                )
                workbook.close()
                screenshot_path = reports.resolve_evidence_path(
                    run.id, data["steps"][0].screenshot_path
                )
                foreign_path = reports.resolve_evidence_path(
                    run.id, data["steps"][1].screenshot_path
                )

            self.assertIn("&lt;script&gt;", html)
            self.assertIn("&lt;img", html)
            self.assertNotIn("<script>alert", html)
            self.assertNotIn("supersecret", html + excel_text)
            self.assertNotIn("abc123", html + excel_text)
            self.assertNotIn("xyz.token", html + excel_text)
            self.assertNotIn("privatetoken", html + excel_text)
            self.assertEqual(html.count("View screenshot"), 1)
            self.assertIn("No screenshot evidence recorded.", html)
            self.assertIsNotNone(screenshot_path)
            self.assertIsNone(foreign_path)


class TestHistoryAndReportApi(unittest.TestCase):
    def test_history_downloads_invalid_ids_and_screenshot_scope(self):
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            original_database_path = database.DATABASE_PATH
            database.DATABASE_PATH = root / "api.db"
            try:
                with (
                    patch.object(reports, "REPORTS_DIR", root / "reports"),
                    patch.object(reports, "PROJECT_DIR", root / "project"),
                    TestClient(app, headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"}) as client,
                ):
                    created = client.post(
                        "/test-runs",
                        json={
                            "run_name": "API history",
                            "platform": "website",
                            "mode": "user-directed",
                            "target": "https://example.test/login",
                            "goal": "Verify rejection.",
                        },
                    )
                    self.assertEqual(created.status_code, 201)
                    run_id = created.json()["id"]

                    history = client.get(
                        "/test-runs?platform=website&limit=1&offset=0"
                    )
                    self.assertEqual(history.status_code, 200)
                    self.assertEqual(history.json()["total"], 1)
                    self.assertEqual(history.json()["items"][0]["id"], run_id)

                    for fmt, media_type in (
                        ("excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                        ("html", "text/html"),
                    ):
                        response = client.get(
                            f"/test-runs/{run_id}/reports/{fmt}"
                        )
                        self.assertEqual(response.status_code, 200)
                        self.assertIn(media_type, response.headers["content-type"])
                        self.assertGreater(len(response.content), 0)

                    run_evidence = root / "project" / "evidence" / run_id
                    run_evidence.mkdir(parents=True)
                    (run_evidence / "step-1.png").write_bytes(b"run screenshot")
                    database.add_test_step(
                        run_id,
                        TestStepCreate(
                            step_number=1,
                            action="Capture screenshot",
                            status="passed",
                            screenshot_path=f"evidence/{run_id}/step-1.png",
                        ),
                        root / "api.db",
                    )
                    screenshot_response = client.get(
                        f"/test-runs/{run_id}/steps/1/screenshot"
                    )
                    self.assertEqual(screenshot_response.status_code, 200)
                    self.assertEqual(screenshot_response.content, b"run screenshot")

                    foreign_id = str(uuid4())
                    foreign_directory = root / "project" / "evidence" / foreign_id
                    foreign_directory.mkdir(parents=True)
                    (foreign_directory / "foreign.png").write_bytes(b"foreign")
                    database.add_test_step(
                        run_id,
                        TestStepCreate(
                            step_number=2,
                            action="Attempt foreign evidence",
                            status="passed",
                            screenshot_path=f"evidence/{foreign_id}/foreign.png",
                        ),
                        root / "api.db",
                    )
                    foreign_response = client.get(
                        f"/test-runs/{run_id}/steps/2/screenshot"
                    )
                    self.assertEqual(foreign_response.status_code, 404)

                    invalid_id = client.get(
                        "/test-runs/not-a-uuid/reports/html"
                    )
                    missing = client.get(
                        f"/test-runs/{uuid4()}/reports/html"
                    )
                    self.assertEqual(invalid_id.status_code, 422)
                    self.assertEqual(missing.status_code, 404)
            finally:
                database.DATABASE_PATH = original_database_path


if __name__ == "__main__":
    unittest.main()
