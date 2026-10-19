import ipaddress
import os
import threading
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from backend import database, reports, website_runner
from backend.reports import generate_excel_report, generate_html_report
from backend.schemas import AutonomousActionPlan, TestRunCreate, TestStepCreate
from backend.website_runner import (
    TargetValidationError,
    _check_expected_outcome_on_page,
    _observe_page,
    _validate_autonomous_plan,
    execute_autonomous_website_run,
    validate_run_target_url,
)

FIXTURE_HOST = "example.com"
FIXTURE_IP = "93.184.216.34"
TARGET = "https://example.com/portal"


def _make_autonomous_run(
    db_path: Path,
    *,
    expected_results: list[str] | None = None,
    rejection_text: str | None = None,
    allow_form_submission: bool = False,
    max_steps: int = 8,
    max_duration_seconds: int = 30,
) -> database.TestRun:
    database.initialize_database(db_path)
    with patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}):
        req = TestRunCreate(
            run_name="Smarter Exploration Test",
            platform="website",
            mode="autonomous-exploration",
            target=TARGET,
            goal="Explore the portal and verify expected outcomes.",
            expected_results=expected_results or [],
            rejection_text=rejection_text,
            allow_form_submission=allow_form_submission,
            max_steps=max_steps,
            max_duration_seconds=max_duration_seconds,
            allowed_hosts=[],
        )
    return database.create_test_run(req, db_path)


class TestSmarterElementDiscovery(unittest.TestCase):
    """Test DOM discovery: accessible labels, toggles, deduplication, and normalized URLs."""

    def test_observe_page_element_discovery_and_deduplication(self):
        html_content = """
        <!doctype html>
        <html>
        <head><title>Discovery Test Page</title></head>
        <body>
            <!-- Duplicate Headings -->
            <h1>Overview</h1>
            <h2>Overview</h2>
            <h3>Unique Section</h3>

            <!-- Links with aria-label, title, fragment, and javascript: -->
            <a href="/dashboard">Dashboard</a>
            <a href="/dashboard" aria-label="Go to Dashboard">Duplicate Link</a>
            <a href="#section-1">Jump to section</a>
            <a href="javascript:void(0)">JS Action</a>
            <a href="/settings" title="User Settings" style="display:inline-block;width:24px;height:24px;">⚙️</a>

            <!-- Buttons: standard, role=button, input[type=button], aria-expanded, safe labels -->
            <button type="button">Save Changes</button>
            <div role="button" tabindex="0" aria-label="Custom Role Button">Custom</div>
            <input type="button" value="Cancel Action">
            <button type="button" aria-expanded="false" aria-controls="faq-1">Toggle FAQ</button>
            <button type="button">Show Details</button>
            <button type="button">Show Details</button> <!-- Duplicate safe button text -->

            <!-- Inputs with label for=id and parent label -->
            <div>
                <label for="search-input">Search Docs</label>
                <input id="search-input" type="search" placeholder="Search...">
            </div>
            <div>
                <label>
                    Filter Category
                    <input type="text" name="category">
                </label>
            </div>
            <!-- Sensitive input should be filtered from safe_to_fill -->
            <input type="password" name="user_password">
        </body>
        </html>
        """
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.route(
                    "https://example.com/test",
                    lambda route: route.fulfill(
                        status=200, content_type="text/html", body=html_content
                    ),
                )
                page.goto("https://example.com/test")

                snapshot, candidates = _observe_page(
                    page,
                    target_host="example.com",
                    allowed_hosts=[],
                    allow_form_submission=False,
                )

                # 1. Heading verification & deduplication in candidates
                heading_candidates = [c for c in candidates.values() if c["kind"] == "heading"]
                heading_texts = {c["text"] for c in heading_candidates}
                self.assertIn("Overview", heading_texts)
                self.assertIn("Unique Section", heading_texts)
                self.assertEqual(len(heading_texts), 2)

                # 2. Link filtering & destination deduplication
                link_candidates = [c for c in candidates.values() if c["kind"] == "link"]
                link_urls = [c["url"] for c in link_candidates]
                # /dashboard should only appear once among candidates
                self.assertEqual(link_urls.count("https://example.com/dashboard"), 1)
                # Fragment '#' and 'javascript:' must be excluded
                for u in link_urls:
                    self.assertFalse(u.endswith("#section-1"))
                    self.assertNotIn("javascript:", u)
                # Title attribute fallback
                settings_candidate = next(
                    (c for c in link_candidates if "settings" in str(c["url"])), None
                )
                self.assertIsNotNone(settings_candidate)
                self.assertEqual(settings_candidate["text"], "User Settings")

                # 3. Button verification: standard, role=button, input[type=button], aria-expanded
                snapshot_btn_texts = {b["text"] for b in snapshot["buttons"]}
                self.assertIn("Save Changes", snapshot_btn_texts)
                self.assertIn("Custom Role Button", snapshot_btn_texts)
                self.assertIn("Cancel Action", snapshot_btn_texts)
                self.assertIn("Toggle FAQ", snapshot_btn_texts)
                self.assertIn("Show Details", snapshot_btn_texts)

                # Safe candidate filtering and deduplication
                button_candidates = [c for c in candidates.values() if c["kind"] == "button"]
                candidate_labels = [c["text"] for c in button_candidates]
                self.assertIn("Toggle FAQ", candidate_labels)
                self.assertIn("Show Details", candidate_labels)
                # Unsafe buttons must not be in candidate action pool
                self.assertNotIn("Save Changes", candidate_labels)
                self.assertNotIn("Cancel Action", candidate_labels)
                # Duplicate safe button "Show Details" must only produce 1 candidate
                self.assertEqual(candidate_labels.count("Show Details"), 1)

                # 4. Input labels association
                input_candidates = [c for c in candidates.values() if c["kind"] == "input"]
                input_labels = [c["text"] for c in input_candidates]
                self.assertTrue(any("Search Docs" in lbl for lbl in input_labels))
                self.assertTrue(any("Filter Category" in lbl for lbl in input_labels))
                # Sensitive password field must not be in safe_to_fill candidates
                self.assertFalse(any("user_password" in c["text"] for c in input_candidates))
            finally:
                browser.close()


class TestAssertionSemanticsAndConflationRegression(unittest.TestCase):
    """Test assertion semantics and regression: unrelated heading checks must not mark goals passed."""

    def test_rejection_regression_unrelated_heading_does_not_pass_configured_outcome(self):
        """Regression Test: An unrelated heading visibility check must never mark

        a configured rejection message or expected result as passed.
        """
        with TemporaryDirectory() as temp_dir:
            db_path = Path(temp_dir) / "test.db"
            run = _make_autonomous_run(
                db_path,
                expected_results=["Account deleted successfully"],
                rejection_text="Invalid credentials",
            )
            database.update_test_run_status(run.id, "running", db_path)

            # Record a step where an arbitrary heading check passed (the conflation defect)
            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="autonomous:assert_visible",
                    page_or_screen="https://example.com/portal",
                    expected_result="Visible element 'heading-0'",
                    actual_result="Element 'heading-0' is visible: Welcome to the Portal",
                    status="passed",
                    duration_ms=100,
                ),
                db_path,
            )
            database.update_test_run_status(run.id, "completed", db_path)

            loaded_run = database.get_test_run(run.id, db_path)
            self.assertEqual(loaded_run.status, "completed")
            # CRITICAL ASSERTION: Must NOT be "passed" - must be "not_evaluated"
            self.assertEqual(loaded_run.assertion_result, "not_evaluated")
            self.assertEqual(loaded_run.suspected_bug_count, 0)
            self.assertIsNone(loaded_run.execution_error)

    def test_expected_result_evaluation_pass_when_verified_on_page(self):
        """When the expected result is actually verified on the page,

        assertion_result becomes 'passed'.
        """
        with TemporaryDirectory() as temp_dir:
            db_path = Path(temp_dir) / "test.db"
            run = _make_autonomous_run(
                db_path,
                expected_results=["Order Confirmed #12345"],
            )
            database.update_test_run_status(run.id, "running", db_path)

            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="autonomous:assert_outcome",
                    page_or_screen="https://example.com/checkout",
                    expected_result="Expected outcome is visible: 'Order Confirmed #12345'",
                    actual_result="Verified outcome 'Order Confirmed #12345' on page: Thank you, Order Confirmed #12345.",
                    status="passed",
                    duration_ms=120,
                ),
                db_path,
            )
            database.update_test_run_status(run.id, "completed", db_path)

            loaded_run = database.get_test_run(run.id, db_path)
            self.assertEqual(loaded_run.status, "completed")
            self.assertEqual(loaded_run.assertion_result, "passed")

    def test_expected_result_failure_marks_run_failed(self):
        """When any required assertion fails, assertion_result becomes 'failed'."""
        with TemporaryDirectory() as temp_dir:
            db_path = Path(temp_dir) / "test.db"
            run = _make_autonomous_run(
                db_path,
                expected_results=["Expected success banner"],
            )
            database.update_test_run_status(run.id, "running", db_path)

            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="autonomous:assert_outcome",
                    page_or_screen="https://example.com/portal",
                    expected_result="Expected outcome is visible: 'Expected success banner'",
                    actual_result="Outcome 'Expected success banner' not found on page.",
                    status="failed",
                    error_details="Expected text was missing.",
                    duration_ms=150,
                ),
                db_path,
            )
            database.update_test_run_status(run.id, "completed", db_path)

            loaded_run = database.get_test_run(run.id, db_path)
            self.assertEqual(loaded_run.status, "completed")
            self.assertEqual(loaded_run.assertion_result, "failed")

    def test_execution_status_separated_from_assertion_result(self):
        """Execution status 'completed' must remain distinct from assertion_result

        'failed' or 'not_evaluated'.
        """
        with TemporaryDirectory() as temp_dir:
            db_path = Path(temp_dir) / "test.db"
            run = _make_autonomous_run(
                db_path,
                expected_results=["Nonexistent element"],
            )
            database.update_test_run_status(run.id, "running", db_path)
            database.update_test_run_status(run.id, "completed", db_path)

            loaded_run = database.get_test_run(run.id, db_path)
            # Run executed to completion without runtime error, but assertions were not evaluated
            self.assertEqual(loaded_run.status, "completed")
            self.assertEqual(loaded_run.assertion_result, "not_evaluated")
            self.assertIsNone(loaded_run.execution_error)


class TestSmarterExplorationExecution(unittest.TestCase):
    """End-to-end autonomous exploration runner tests with mocked Playwright interception."""

    def test_autonomous_run_evaluates_expected_outcome_on_page(self):
        """Runner actively evaluates required expected results against page content."""
        html_page = """
        <!doctype html>
        <html>
        <head><title>Portal Test</title></head>
        <body>
            <h1>Customer Portal</h1>
            <div id="notice" class="alert-success">System operational and ready.</div>
            <a href="/portal/subpage">Subpage</a>
        </body>
        </html>
        """
        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(
                database.DATABASE_PATH,
                expected_results=["System operational and ready."],
            )
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            def intercept_page(route):
                route.fulfill(status=200, content_type="text/html", body=html_page)
                return True

            cancel_event = threading.Event()
            # Single-step planner that inspects, allowing page outcome evaluation
            plans = iter(
                [
                    '{"actions":[{"action":"inspect","element_id":"heading-0"}]}',
                ]
            )

            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=lambda prompt: next(plans, ""),
                    request_interceptor=intercept_page,
                )

            final_run = database.get_test_run(run.id)
            steps = database.list_test_steps(run.id)
            self.assertEqual(final_run.status, "completed")
            self.assertEqual(final_run.assertion_result, "passed")

            # Check that an autonomous:assert_outcome step was logged
            outcome_steps = [s for s in steps if s.action == "autonomous:assert_outcome"]
            self.assertGreaterEqual(len(outcome_steps), 1)
            self.assertEqual(outcome_steps[0].status, "passed")
            self.assertIn("System operational and ready.", outcome_steps[0].expected_result)
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()

    def test_planner_fallback_handles_malformed_plan_without_crashing(self):
        """When planner throws or generates invalid JSON, fallback is safely invoked."""
        html_page = """
        <!doctype html>
        <html>
        <head><title>Fallback Portal</title></head>
        <body>
            <h1>Portal Header</h1>
            <a href="/portal/about">About us</a>
        </body>
        </html>
        """
        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(
                database.DATABASE_PATH,
                expected_results=["Expected text not on page"],
            )
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            def intercept_page(route):
                route.fulfill(status=200, content_type="text/html", body=html_page)
                return True

            cancel_event = threading.Event()

            # Planner returns garbage JSON
            def bad_planner(prompt):
                return "THIS IS NOT VALID JSON"

            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=bad_planner,
                    request_interceptor=intercept_page,
                )

            final_run = database.get_test_run(run.id)
            report_data = database.get_test_run_report_data(run.id)
            self.assertEqual(final_run.status, "completed")
            # Unrelated fallback action must not mark missing expected outcome as passed!
            self.assertEqual(final_run.assertion_result, "not_evaluated")
            # Verify planning_fallback log was recorded
            self.assertTrue(
                any(log.event_type == "planning_fallback" for log in report_data["logs"])
            )
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()

    def test_cancellation_stops_autonomous_run_cleanly(self):
        """Setting cancel_event stops exploration promptly."""
        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(database.DATABASE_PATH, max_steps=20)
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            cancel_event = threading.Event()
            # Cancel immediately
            cancel_event.set()

            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=lambda p: '{"actions":[{"action":"inspect","element_id":"heading-0"}]}',
                    request_interceptor=lambda route: route.fulfill(
                        status=200, content_type="text/html", body="<h1>Test</h1>"
                    ),
                )

            final_run = database.get_test_run(run.id)
            self.assertEqual(final_run.status, "cancelled")
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()

    def test_browser_timeout_terminates_run_and_records_failed_status(self):
        """Playwright TimeoutError terminates exploration immediately and records failed status with execution_error."""
        from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError

        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(database.DATABASE_PATH, max_steps=10)
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            cancel_event = threading.Event()

            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
                patch.object(
                    Page,
                    "goto",
                    side_effect=PlaywrightTimeoutError("Navigation timeout of 5000ms exceeded."),
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=lambda p: '{"actions":[{"action":"inspect","element_id":"heading-0"}]}',
                )

            final_run = database.get_test_run(run.id)
            report_data = database.get_test_run_report_data(run.id)
            steps = report_data["steps"]

            # Must be marked failed, not completed
            self.assertEqual(final_run.status, "failed")
            self.assertIn("Website action failed (TimeoutError).", final_run.error_details or "")
            self.assertEqual(final_run.execution_error, "browser")
            self.assertEqual(final_run.assertion_result, "not_evaluated")

            # Must stop immediately after step 1 error without attempting further actions
            self.assertEqual(len(steps), 1)
            self.assertEqual(steps[0].step_number, 1)
            self.assertEqual(steps[0].status, "error")
            self.assertIn("Browser Error: TimeoutError.", steps[0].actual_result or "")

            # Verify browser_error log was recorded
            self.assertTrue(
                any(log.event_type == "browser_error" for log in report_data["logs"])
            )
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()

    def test_autonomous_navigation_uses_dedicated_navigation_timeout(self):
        """Autonomous exploration uses dedicated navigation timeout (up to 15s) for page navigations rather than 5s action timeout."""
        from playwright.sync_api import Page

        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(database.DATABASE_PATH, max_steps=5, max_duration_seconds=30)
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            captured_timeouts: list[int] = []
            orig_goto = Page.goto

            def spy_goto(self, *args, **kwargs):
                if "timeout" in kwargs:
                    captured_timeouts.append(kwargs["timeout"])
                return orig_goto(self, *args, **kwargs)

            cancel_event = threading.Event()

            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
                patch.object(Page, "goto", spy_goto),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=None,
                    request_interceptor=lambda route: route.fulfill(
                        status=200, content_type="text/html", body="<h1>Portal</h1>"
                    ),
                )

            final_run = database.get_test_run(run.id)
            report_data = database.get_test_run_report_data(run.id)

            # Navigation timeout must have been recorded above 5000ms (15000ms)
            self.assertTrue(len(captured_timeouts) >= 1)
            self.assertEqual(captured_timeouts[0], website_runner._MAX_AUTONOMOUS_NAVIGATION_MS)
            self.assertGreater(captured_timeouts[0], website_runner._MAX_AUTONOMOUS_ACTION_MS)

            # Step 1 inspect completed successfully
            self.assertEqual(final_run.status, "completed")
            self.assertEqual(report_data["steps"][0].status, "passed")
            self.assertEqual(report_data["steps"][0].action, "autonomous:inspect")
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()

    def test_security_ssrf_and_navigation_guardrails(self):
        """Target validation blocks private, loopback, and metadata IPs."""
        for target in (
            "http://127.0.0.1/admin",
            "http://169.254.169.254/latest/meta-data/",
            "http://10.0.0.1/secret",
            "http://192.168.1.1/router",
        ):
            with self.subTest(target=target), self.assertRaises(TargetValidationError):
                validate_run_target_url(target, FIXTURE_HOST, [])


class TestSmarterExplorationReports(unittest.TestCase):
    """HTML and Excel report compatibility for outcome assertions."""

    def test_reports_format_outcome_assertions_correctly(self):
        with TemporaryDirectory() as temp_dir:
            db_path = Path(temp_dir) / "reports.db"
            run = _make_autonomous_run(
                db_path,
                expected_results=["Welcome to Dashboard"],
            )
            database.update_test_run_status(run.id, "running", db_path)
            database.add_test_step(
                run.id,
                TestStepCreate(
                    step_number=1,
                    action="autonomous:assert_outcome",
                    page_or_screen="https://example.com/portal",
                    expected_result="Expected outcome is visible: 'Welcome to Dashboard'",
                    actual_result="Verified outcome 'Welcome to Dashboard' on page: Welcome to Dashboard.",
                    status="passed",
                    duration_ms=80,
                ),
                db_path,
            )
            database.update_test_run_status(run.id, "completed", db_path)

            report_data = database.get_test_run_report_data(run.id, db_path)

            # 1. HTML Report
            html_path = generate_html_report(report_data, base_url="http://localhost:8000")
            html_content = html_path.read_text(encoding="utf-8")
            self.assertIn("autonomous:assert_outcome", html_content)
            self.assertIn("Welcome to Dashboard", html_content)
            self.assertIn("Passed", html_content)

            # 2. Excel Report
            excel_path = generate_excel_report(report_data)
            wb = load_workbook(excel_path)
            self.assertIn("Test Steps", wb.sheetnames)
            ws = wb["Test Steps"]
            step_rows = list(ws.iter_rows(values_only=True))
            action_col = [row[1] for row in step_rows if row and len(row) > 1]
            self.assertTrue(any("autonomous:assert_outcome" in str(a) for a in action_col))
            self.assertTrue(any("Passed" in str(row[6]) for row in step_rows if row and len(row) > 6))


class TestZeroElementsAndDiagnosticFindings(unittest.TestCase):
    """Test diagnostic findings when 0 elements or 0 actionable candidates are discovered."""

    def test_react_native_web_element_discovery_and_probe_wait(self):
        """React Native Web elements (tabindex=0, h5/h6, role=heading, tel input) are discovered."""
        html_content = """
        <!doctype html>
        <html>
        <head><title></title></head>
        <body>
            <div role="heading" aria-level="1">MEATZAAR</div>
            <h5>Subtitle Section</h5>
            <input type="tel" placeholder="Phone number">
            <input type="password" placeholder="Password">
            <div tabindex="0">SIGN IN</div>
            <div tabindex="0">CREATE AN ACCOUNT</div>
            <div tabindex="0" role="button">Show Details</div>
            <!-- Container with tabindex="0" containing child controls must be ignored as a button -->
            <div tabindex="0">
                <input type="text" name="child">
            </div>
        </body>
        </html>
        """
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.route(
                    "https://example.com/rnw",
                    lambda route: route.fulfill(
                        status=200, content_type="text/html", body=html_content
                    ),
                )
                page.goto("https://example.com/rnw")

                snapshot, candidates = _observe_page(
                    page,
                    target_host="example.com",
                    allowed_hosts=[],
                    allow_form_submission=False,
                )

                # Headings include role=heading and h5
                heading_texts = {h["text"] for h in snapshot["headings"]}
                self.assertIn("MEATZAAR", heading_texts)
                self.assertIn("Subtitle Section", heading_texts)

                # Buttons include tabindex=0 divs
                btn_texts = {b["text"] for b in snapshot["buttons"]}
                self.assertIn("SIGN IN", btn_texts)
                self.assertIn("CREATE AN ACCOUNT", btn_texts)
                self.assertIn("Show Details", btn_texts)

                # Unsafe buttons not in candidates; safe button in candidates
                safe_btn_candidates = [
                    c for c in candidates.values() if c["kind"] == "button"
                ]
                safe_labels = {c["text"] for c in safe_btn_candidates}
                self.assertIn("Show Details", safe_labels)
                self.assertNotIn("SIGN IN", safe_labels)
                self.assertNotIn("CREATE AN ACCOUNT", safe_labels)

                # Inputs include tel and password
                input_types = {inp["type"] for inp in snapshot["inputs"]}
                self.assertIn("tel", input_types)
                self.assertIn("password", input_types)
            finally:
                browser.close()

    def test_zero_elements_reports_diagnostic_finding_and_fails_run(self):
        """When 0 page elements are found, step 1 fails, logs diagnostic finding, and run fails."""
        empty_html = "<!doctype html><html><head><title>Empty Page</title></head><body></body></html>"
        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(
                database.DATABASE_PATH,
                expected_results=["Some outcome"],
            )
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            cancel_event = threading.Event()
            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=None,
                    request_interceptor=lambda route: route.fulfill(
                        status=200, content_type="text/html", body=empty_html
                    ),
                )

            final_run = database.get_test_run(run.id)
            report_data = database.get_test_run_report_data(run.id)

            # Critical Requirement: Must NOT be completed
            self.assertEqual(final_run.status, "failed")
            self.assertEqual(final_run.execution_error, "browser")
            self.assertIn("Diagnostic finding", final_run.error_details)
            self.assertEqual(final_run.assertion_result, "not_evaluated")

            # Step 1 inspect must have failed and recorded diagnostic finding
            self.assertEqual(len(report_data["steps"]), 1)
            step1 = report_data["steps"][0]
            self.assertEqual(step1.status, "failed")
            self.assertIn("0 headings, 0 links", step1.actual_result)
            self.assertIn("Diagnostic finding", step1.actual_result)

            # Execution logs must include diagnostic_finding event
            self.assertTrue(
                any(log.event_type == "diagnostic_finding" for log in report_data["logs"])
            )

            # Bug record for zero page elements must be logged
            bugs = database.list_bug_records(run.id)
            self.assertEqual(len(bugs), 1)
            self.assertIn("zero page elements", bugs[0].title.lower())
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()

    def test_read_only_login_page_with_no_safe_actions_completes_with_diagnostic_finding(self):
        """When page has elements (e.g. login form) but 0 safe actions, read-only inspection completes successfully."""
        login_html = """
        <!doctype html>
        <html>
        <head><title>Login Page</title></head>
        <body>
            <input type="tel" name="phone" placeholder="Phone">
            <input type="password" name="password" placeholder="Password">
            <button type="submit">Sign In</button>
        </body>
        </html>
        """
        temp_dir = TemporaryDirectory()
        root = Path(temp_dir.name)
        old_db_path = database.DATABASE_PATH
        old_evidence = website_runner.EVIDENCE_DIR
        old_project = website_runner.PROJECT_DIR

        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root

        try:
            run = _make_autonomous_run(
                database.DATABASE_PATH,
                expected_results=["Dashboard"],
            )
            database.update_test_run_status(run.id, "running")
            self.assertTrue(website_runner.try_reserve_run())

            cancel_event = threading.Event()
            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch(
                    "backend.website_runner.validate_target_url",
                    return_value=FIXTURE_HOST,
                ),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event,
                    planner=None,
                    request_interceptor=lambda route: route.fulfill(
                        status=200, content_type="text/html", body=login_html
                    ),
                )

            final_run = database.get_test_run(run.id)
            report_data = database.get_test_run_report_data(run.id)

            # Read-only inspection must succeed: completed status, no execution error
            self.assertEqual(final_run.status, "completed")
            self.assertIsNone(final_run.execution_error)
            self.assertIsNone(final_run.error_details)
            # Semantics accurate: no assertions evaluated must remain not_evaluated, not passed
            self.assertEqual(final_run.assertion_result, "not_evaluated")

            # Step 1 passed because elements were observed
            self.assertEqual(len(report_data["steps"]), 1)
            step1 = report_data["steps"][0]
            self.assertEqual(step1.status, "passed")
            self.assertIn("2 inputs", step1.actual_result)
            self.assertIn("0 safe search", step1.actual_result)

            # Diagnostic finding log recorded without treating finding as a website bug
            self.assertTrue(
                any(log.event_type == "diagnostic_finding" for log in report_data["logs"])
            )
            self.assertTrue(
                any(log.event_type == "exploration_complete" for log in report_data["logs"])
            )
            bugs = database.list_bug_records(run.id)
            self.assertEqual(len(bugs), 0)
        finally:
            database.DATABASE_PATH = old_db_path
            website_runner.EVIDENCE_DIR = old_evidence
            website_runner.PROJECT_DIR = old_project
            temp_dir.cleanup()


if __name__ == "__main__":
    unittest.main()
