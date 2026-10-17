import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import database, reports
from backend.app import app
from backend.schemas import BugRecordCreate, TestRunCreate, TestStepCreate
from backend.security import API_TOKEN


def make_run(path: Path, name: str, status: str = "pending"):
    database.initialize_database(path)
    run = database.create_test_run(
        TestRunCreate(
            run_name=name,
            platform="website",
            mode="user-directed",
            target="https://example.test/login",
            goal="Check a saved result.",
        ),
        path,
    )
    if status != "pending":
        database.update_test_run_status(run.id, "running", path)
        database.update_test_run_status(run.id, status, path)
    return run


class TestDashboardAccuracy(unittest.TestCase):
    def test_empty_metrics_and_history_are_zero_and_empty(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "empty.db"
            database.initialize_database(path)
            self.assertEqual(
                database.get_dashboard_metrics(path),
                {
                    "total_runs": 0,
                    "completed_runs": 0,
                    "failed_runs": 0,
                    "cancelled_runs": 0,
                    "assertions_passed": 0,
                    "assertions_failed": 0,
                    "suspected_bugs": 0,
                },
            )
            self.assertEqual(database.list_test_run_history(database_path=path), ([], 0))

    def test_metrics_distinguish_execution_and_assertion_outcomes(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "metrics.db"
            passed = make_run(path, "passed", "completed")
            failed = make_run(path, "failed", "completed")
            cancelled = make_run(path, "cancelled", "cancelled")
            pending = make_run(path, "pending")
            database.add_test_step(
                passed.id,
                TestStepCreate(step_number=1, action="check_visibility", status="passed"),
                path,
            )
            database.add_test_step(
                failed.id,
                TestStepCreate(step_number=1, action="check_visibility", status="failed"),
                path,
            )
            database.add_bug_record(
                failed.id,
                BugRecordCreate(
                    title="Suspected defect",
                    description="The assertion failed.",
                    severity="Medium",
                    confidence=0.8,
                ),
                path,
            )
            self.assertEqual(database.get_test_run(passed.id, path).assertion_result, "passed")
            self.assertEqual(database.get_test_run(failed.id, path).assertion_result, "failed")
            self.assertEqual(
                database.get_dashboard_metrics(path),
                {
                    "total_runs": 4,
                    "completed_runs": 2,
                    "failed_runs": 0,
                    "cancelled_runs": 1,
                    "assertions_passed": 1,
                    "assertions_failed": 1,
                    "suspected_bugs": 1,
                },
            )
            self.assertEqual(pending.status, "pending")
            self.assertEqual(database.get_test_run(cancelled.id, path).status, "cancelled")

    def test_reports_and_api_use_the_same_derived_values(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "report.db"
            run = make_run(path, "report", "completed")
            database.add_test_step(
                run.id,
                TestStepCreate(step_number=1, action="inspect_page", status="passed"),
                path,
            )
            database.add_test_step(
                run.id,
                TestStepCreate(step_number=2, action="check_visibility", status="failed"),
                path,
            )
            data = database.get_test_run_report_data(run.id, path)
            old_reports, old_project = reports.REPORTS_DIR, reports.PROJECT_DIR
            reports.REPORTS_DIR = root / "reports"
            reports.PROJECT_DIR = root
            try:
                html = reports.generate_html_report(data, "http://127.0.0.1:8000")
                workbook = load_workbook(reports.generate_excel_report(data), read_only=True)
                summary = "\n".join(
                    f"{row[0]}:{row[1]}"
                    for row in workbook["Test Summary"].iter_rows(values_only=True)
                )
                workbook.close()
            finally:
                reports.REPORTS_DIR, reports.PROJECT_DIR = old_reports, old_project
            self.assertIn("Assertions passed:0", summary)
            self.assertIn("Assertions failed:1", summary)
            html_text = html.read_text(encoding="utf-8")
            self.assertIn("Assertion outcome:</strong> Not evaluated", html_text)
            self.assertIn("Assertion outcome:</strong> Failed", html_text)

    def test_dashboard_endpoint_requires_authentication(self):
        with TestClient(app, headers={"Host": "127.0.0.1:8000"}) as client:
            self.assertEqual(client.get("/dashboard/metrics").status_code, 401)
            response = client.get(
                "/dashboard/metrics",
                headers={"Authorization": f"Bearer {API_TOKEN}"},
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(set(response.json()), {
                "total_runs",
                "completed_runs",
                "failed_runs",
                "cancelled_runs",
                "assertions_passed",
                "assertions_failed",
                "suspected_bugs",
            })


if __name__ == "__main__":
    unittest.main()
