import ipaddress
import os
import threading
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

os.environ.setdefault("GEMINI_API_KEY", "milestone-test-key")

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from backend import database, reports, website_runner
from backend.app import app
from backend.security import API_TOKEN
from backend.schemas import AutonomousActionPlan, TestRunCreate
from backend.website_runner import (
    TargetValidationError,
    _validate_autonomous_plan,
    execute_autonomous_website_run,
    validate_allowed_hosts,
    validate_run_target_url,
)


TARGET = "https://example.com/start"
FIXTURE_HOST = "example.com"
FIXTURE_IP = "93.184.216.34"
FIXTURE_REQUESTS: list[str] = []
FORM_MUTATION_SCRIPT = ""


def make_autonomous_run(
    database_path: Path,
    *,
    max_steps: int = 12,
    max_duration_seconds: int = 30,
    allowed_hosts: list[str] | None = None,
    allow_form_submission: bool = False,
):
    database.initialize_database(database_path)
    with patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}):
        request = TestRunCreate(
            run_name="Autonomous fixture exploration",
            platform="website",
            mode="autonomous-exploration",
            target=TARGET,
            goal="Inspect the fixture and safely examine its search and documentation.",
            max_steps=max_steps,
            max_duration_seconds=max_duration_seconds,
            allowed_hosts=allowed_hosts or [],
            allow_form_submission=allow_form_submission,
        )
    return database.create_test_run(request, database_path)


def fixture_response(url: str) -> str:
    parsed = website_runner.urlsplit(url)
    if parsed.path == "/next":
        return """
            <!doctype html>
            <html><head><title>Next fixture page</title></head>
            <body><h1>Next page</h1><p>Safe destination content.</p></body></html>
        """
    if parsed.path.startswith("/chain/"):
        try:
            index = int(parsed.path.rsplit("/", 1)[-1])
        except ValueError:
            index = 0
        return (
            "<!doctype html><html><head><title>Navigation fixture</title></head><body>"
            f"<h1>Page {index}</h1>"
            f'<a href="/chain/{index + 1}">Page {index + 1}</a>'
            "</body></html>"
        )
    if parsed.path == "/search":
        return """
            <!doctype html>
            <html><head><title>Search results</title></head>
            <body><h1>Search results</h1><p>Results for the safe query.</p></body></html>
        """
    return f"""
        <!doctype html>
        <html><head><title>QAI safe fixture</title></head>
        <body>
          <h1>Fixture home</h1>
          <input id="search" type="search" placeholder="Search catalogue">
          <form action="/search" method="get">
            <input type="text" name="query" placeholder="Filter catalogue">
            <button type="submit">Apply filter</button>
          </form>
          <a href="/next">Safe documentation</a>
          <a href="http://169.254.169.254/latest/meta-data/">Metadata</a>
          <form action="/must-not-submit">
            <input type="text" name="username">
            <input type="password" name="password">
            <button type="submit">Delete account</button>
          </form>
          <button type="button">Delete everything</button>
          <script>{FORM_MUTATION_SCRIPT}</script>
        </body></html>
    """


def intercept_fixture(route):
    url = route.request.url
    FIXTURE_REQUESTS.append(url)
    if url.startswith("https://example.com/"):
        route.fulfill(
            status=200,
            content_type="text/html",
            body=fixture_response(url),
        )
    else:
        route.abort()
    return True


class TestAutonomousActionValidation(unittest.TestCase):
    def test_form_submission_defaults_to_denied_and_round_trips_opt_in(self):
        with TemporaryDirectory() as temporary_directory:
            database_path = Path(temporary_directory) / "authorization.db"
            denied = make_autonomous_run(database_path)
            self.assertFalse(denied.allow_form_submission)
            self.assertFalse(
                database.get_test_run(denied.id, database_path).allow_form_submission
            )

            allowed = make_autonomous_run(
                database_path,
                allow_form_submission=True,
            )
            self.assertTrue(
                database.get_test_run(allowed.id, database_path).allow_form_submission
            )

    def test_strict_plan_validation_and_observed_candidate_binding(self):
        valid = AutonomousActionPlan.model_validate_json(
            '{"actions":[{"action":"fill","element_id":"input-0","value":"catalog"}]}'
        )
        candidates = {
            "input-0": {"kind": "input"},
            "link-0": {"kind": "link", "url": "https://example.com/next"},
            "button-0": {"kind": "button"},
        }
        self.assertEqual(valid.actions[0].action, "fill")
        self.assertIsNotNone(
            _validate_autonomous_plan(
                '{"actions":[{"action":"fill","element_id":"input-0","value":"catalog"}]}',
                candidates,
            )
        )
        self.assertIsNotNone(
            _validate_autonomous_plan(
                '{"actions":[{"action":"submit_form","element_id":"form-0"}]}',
                {
                    **candidates,
                    "form-0": {"kind": "form", "safe_to_submit": True},
                },
            )
        )
        self.assertIsNone(
            _validate_autonomous_plan(
                '{"actions":[{"action":"click","element_id":"input-0"}]}',
                candidates,
            )
        )
        self.assertIsNone(
            _validate_autonomous_plan(
                '{"actions":[{"action":"navigate","url":"http://169.254.169.254/"}]}',
                candidates,
            )
        )
        for raw in (
            '{"actions":[{"action":"submit_form","selector":"form"}]}',
            '{"actions":[{"action":"click","element_id":"button-0","code":"alert(1)"}]}',
            '{"actions":[{"action":"inspect","element_id":"link-0"}]}',
            '{"actions":[{"action":"fill","element_id":"input-0","value":""}]}',
        ):
            with self.subTest(raw=raw), self.assertRaises(ValidationError):
                AutonomousActionPlan.model_validate_json(raw)

    def test_form_submission_requires_an_authorized_observed_candidate(self):
        self.assertIsNone(
            _validate_autonomous_plan(
                '{"actions":[{"action":"submit_form","element_id":"form-0"}]}',
                {"form-0": {"kind": "form", "safe_to_submit": False}},
            )
        )

    def test_target_and_allowlist_reject_unsafe_hosts(self):
        self.assertEqual(validate_allowed_hosts(["Docs.Example.com"]), ["docs.example.com"])
        for host in ("127.0.0.1", "169.254.169.254", "*.example.com", "example.com:443"):
            with self.subTest(host=host), self.assertRaises(ValueError):
                validate_allowed_hosts([host])
        with self.assertRaises(TargetValidationError):
            validate_run_target_url(
                "http://169.254.169.254/latest/meta-data/",
                FIXTURE_HOST,
                [],
            )
        with self.assertRaises(TargetValidationError):
            validate_run_target_url(
                "https://unapproved.example/",
                FIXTURE_HOST,
                [],
                resolve_host=False,
            )

    def test_autonomous_request_limits_and_mode_round_trip(self):
        with TemporaryDirectory() as temporary_directory:
            database_path = Path(temporary_directory) / "mode.db"
            run = make_autonomous_run(
                database_path,
                max_steps=12,
                allowed_hosts=["docs.example.com"],
            )
            loaded = database.get_test_run(run.id, database_path)
            self.assertEqual(run.mode, "autonomous-exploration")
            self.assertEqual(loaded.mode, "autonomous-exploration")
            self.assertEqual(loaded.allowed_hosts, ["docs.example.com"])

            with database._connection(database_path) as connection:
                connection.execute(
                    """
                    UPDATE test_runs
                    SET mode = 'autonomous', autonomous_exploration = 0
                    WHERE id = ?
                    """,
                    (run.id,),
                )
            legacy = database.get_test_run(run.id, database_path)
            history, _ = database.list_test_run_history(database_path=database_path)
            self.assertEqual(legacy.mode, "autonomous")
            self.assertEqual(history[0].mode, "autonomous")

        with self.assertRaises(ValidationError):
            TestRunCreate(
                run_name="Too many steps",
                platform="website",
                mode="autonomous-exploration",
                target=TARGET,
                goal="Try to run too long.",
                max_steps=51,
            )
        with self.assertRaises(ValidationError):
            TestRunCreate(
                run_name="Too much time",
                platform="website",
                mode="autonomous-exploration",
                target=TARGET,
                goal="Try to run too long.",
                max_duration_seconds=601,
            )


class TestAutonomousPlaywrightExecution(unittest.TestCase):
    def test_authorized_safe_form_is_filled_submitted_and_verified(self):
        FIXTURE_REQUESTS.clear()
        plans = iter(
            [
                '{"actions":[{"action":"fill","element_id":"input-1","value":"catalog"}]}',
                '{"actions":[{"action":"submit_form","element_id":"form-0"}]}',
                '{"actions":[{"action":"assert_visible","element_id":"heading-0","expected_text":"Search results"}]}',
            ]
        )
        result = self.execute_fixture(
            lambda prompt: next(plans),
            allow_form_submission=True,
        )
        _, run, steps, bugs, _, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertEqual(
                [step.action for step in steps[:3]],
                [
                    "autonomous:inspect",
                    "autonomous:fill",
                    "autonomous:submit_form",
                ],
            )
            self.assertTrue(any(step.action == "autonomous:assert_visible" for step in steps))
            self.assertTrue(any("/search" in url for url in FIXTURE_REQUESTS))
            self.assertEqual(bugs, [])
        finally:
            temporary_directory.cleanup()

    def test_authorized_run_rejects_destructive_form_submission(self):
        FIXTURE_REQUESTS.clear()
        result = self.execute_fixture(
            lambda prompt: '{"actions":[{"action":"submit_form","element_id":"form-1"}]}',
            allow_form_submission=True,
        )
        _, run, steps, _, _, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertFalse(any("/must-not-submit" in url for url in FIXTURE_REQUESTS))
        finally:
            temporary_directory.cleanup()

    def test_live_form_method_mutation_is_blocked(self):
        self.assert_live_form_mutation_blocked(
            "document.querySelector('form').method = 'post';"
        )

    def test_live_form_destination_mutation_is_blocked(self):
        self.assert_live_form_mutation_blocked(
            "document.querySelector('form').action = 'http://169.254.169.254/';"
        )

    def test_live_form_sensitive_field_mutation_is_blocked(self):
        self.assert_live_form_mutation_blocked(
            "document.querySelector('form').insertAdjacentHTML("
            "'beforeend', '<input name=\"password\" type=\"password\">');"
        )

    def test_live_form_submit_control_replacement_is_blocked(self):
        self.assert_live_form_mutation_blocked(
            "document.querySelector('form button[type=\"submit\"]').replaceWith("
            "'<button type=\"submit\">Apply filter</button>');"
        )

    def test_live_form_disabled_submit_control_is_blocked(self):
        self.assert_live_form_mutation_blocked(
            "document.querySelector('form button[type=\"submit\"]').disabled = true;"
        )

    def assert_live_form_mutation_blocked(self, mutation_script: str):
        global FORM_MUTATION_SCRIPT
        FIXTURE_REQUESTS.clear()
        previous_script = FORM_MUTATION_SCRIPT
        FORM_MUTATION_SCRIPT = (
            "document.querySelector('form input[name=\"query\"]').addEventListener("
            "'input', () => {"
            f"{mutation_script}"
            "}, {once: true});"
        )
        try:
            plans = iter(
                [
                    '{"actions":[{"action":"fill","element_id":"input-1","value":"catalog"}]}',
                    '{"actions":[{"action":"submit_form","element_id":"form-0"}]}',
                ]
            )
            result = self.execute_fixture(
                lambda prompt: next(plans),
                allow_form_submission=True,
            )
            _, run, steps, _, _, _, temporary_directory = result
            try:
                self.assertEqual(run.status, "completed")
                self.assertFalse(any("/search" in url for url in FIXTURE_REQUESTS))
                self.assertFalse(
                    any(step.action == "autonomous:submit_form" for step in steps)
                )
            finally:
                temporary_directory.cleanup()
        finally:
            FORM_MUTATION_SCRIPT = previous_script

    def test_form_submission_is_denied_by_default(self):
        FIXTURE_REQUESTS.clear()
        result = self.execute_fixture(
            lambda prompt: '{"actions":[{"action":"submit_form","element_id":"form-0"}]}'
        )
        _, run, steps, _, report_data, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertFalse(any(step.action == "autonomous:submit_form" for step in steps))
            self.assertFalse(any("/search" in url for url in FIXTURE_REQUESTS))
            self.assertTrue(
                any(log.event_type == "planning_fallback" for log in report_data["logs"])
            )
        finally:
            temporary_directory.cleanup()

    def execute_fixture(
        self,
        planner,
        *,
        max_steps: int = 12,
        max_duration_seconds: int = 30,
        initial_url: str = TARGET,
        cancel_event: threading.Event | None = None,
        allow_form_submission: bool = False,
    ):
        temporary_directory = TemporaryDirectory()
        root = Path(temporary_directory.name)
        old_database_path = database.DATABASE_PATH
        old_evidence_dir = website_runner.EVIDENCE_DIR
        old_project_dir = website_runner.PROJECT_DIR
        database.DATABASE_PATH = root / "runs.db"
        website_runner.EVIDENCE_DIR = root / "evidence"
        website_runner.PROJECT_DIR = root
        run = make_autonomous_run(
            database.DATABASE_PATH,
            max_steps=max_steps,
            max_duration_seconds=max_duration_seconds,
            allow_form_submission=allow_form_submission,
        )
        if initial_url != TARGET:
            run = database.get_test_run(run.id)
            run.target = initial_url
            with database._connection(database.DATABASE_PATH) as connection:
                connection.execute(
                    "UPDATE test_runs SET target = ? WHERE id = ?",
                    (initial_url, run.id),
                )
        database.update_test_run_status(run.id, "running")
        self.assertTrue(website_runner.try_reserve_run())

        def fake_validator(url: str, original_host: str | None = None):
            if website_runner.urlsplit(url).scheme not in {"https", "http"}:
                raise TargetValidationError("Only HTTP and HTTPS allowed.")
            return FIXTURE_HOST

        try:
            with (
                patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                patch("backend.website_runner.validate_target_url", side_effect=fake_validator),
                patch(
                    "backend.website_runner._resolve_public_addresses",
                    return_value={ipaddress.ip_address(FIXTURE_IP)},
                ),
            ):
                execute_autonomous_website_run(
                    run.id,
                    cancel_event or threading.Event(),
                    planner=planner,
                    request_interceptor=intercept_fixture,
                )
            return (
                run.id,
                database.get_test_run(run.id),
                database.list_test_steps(run.id),
                database.list_bug_records(run.id),
                database.get_test_run_report_data(run.id),
                root,
                temporary_directory,
            )
        finally:
            database.DATABASE_PATH = old_database_path
            website_runner.EVIDENCE_DIR = old_evidence_dir
            website_runner.PROJECT_DIR = old_project_dir

    def test_ai_actions_execute_in_real_chromium_and_persist_report_data(self):
        plans = iter(
            [
                '{"actions":[{"action":"fill","element_id":"input-0","value":"catalog"}]}',
                '{"actions":[{"action":"click","element_id":"link-0"}]}',
                '{"actions":[{"action":"assert_visible","element_id":"heading-0","expected_text":"Next page"}]}',
                '{"actions":[{"action":"inspect"}]}',
            ]
        )
        result = self.execute_fixture(lambda prompt: next(plans))
        run_id, run, steps, bugs, report_data, root, temporary_directory = result
        self.assertEqual(run.status, "completed")
        self.assertEqual(run.mode, "autonomous-exploration")
        self.assertEqual([step.status for step in steps], ["passed"] * 4)
        self.assertEqual(steps[1].action, "autonomous:fill")
        self.assertIn("entered text was not logged", steps[1].actual_result)
        self.assertEqual(steps[2].action, "autonomous:click")
        self.assertIn("/next", steps[2].page_or_screen)
        self.assertEqual(steps[3].status, "passed")
        self.assertEqual(bugs, [])
        self.assertTrue(all(step.screenshot_path for step in steps))
        self.assertTrue(
            all((root / step.screenshot_path).is_file() for step in steps)
        )

        report_root = root / "reports"
        old_reports = reports.REPORTS_DIR
        old_report_project = reports.PROJECT_DIR
        try:
            reports.REPORTS_DIR = report_root
            reports.PROJECT_DIR = root
            excel = reports.generate_excel_report(report_data)
            html = reports.generate_html_report(report_data, "http://127.0.0.1:8000")
            workbook = load_workbook(excel, read_only=True)
            summary = list(workbook["Test Summary"].values)
            self.assertIn(("Mode", "autonomous-exploration"), summary)
            self.assertEqual(workbook["Test Steps"].max_row, 5)
            headers = list(next(workbook["Test Steps"].iter_rows(values_only=True)))
            self.assertIn("Evidence availability", headers)
            self.assertIn("Evidence reference", headers)
            workbook.close()
            html_text = html.read_text(encoding="utf-8")
            self.assertIn("autonomous-exploration", html_text)
            self.assertIn("Clicked safe observed link", html_text)
            self.assertIn("Safe documentation", html_text)
            self.assertIn("View screenshot", html_text)
        finally:
            reports.REPORTS_DIR = old_reports
            reports.PROJECT_DIR = old_report_project
            temporary_directory.cleanup()

    def test_screenshot_failure_does_not_fail_successful_action(self):
        with patch.object(
            website_runner.Page,
            "screenshot",
            side_effect=RuntimeError("screenshot unavailable"),
        ):
            result = self.execute_fixture(
                lambda prompt: '{"actions":[{"action":"inspect"}]}',
                max_steps=1,
            )
        _, run, steps, _, report_data, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertEqual(steps[0].status, "passed")
            self.assertIsNone(steps[0].screenshot_path)
            self.assertTrue(
                any(log.event_type == "evidence_error" for log in report_data["logs"])
            )
        finally:
            temporary_directory.cleanup()

    def test_report_rejects_traversal_evidence_and_keeps_reports_consistent(self):
        result = self.execute_fixture(
            lambda prompt: '{"actions":[{"action":"inspect"}]}',
            max_steps=1,
        )
        _, _, steps, _, report_data, root, temporary_directory = result
        try:
            steps[0].screenshot_path = "../outside.png"
            old_reports = reports.REPORTS_DIR
            try:
                reports.REPORTS_DIR = root / "reports"
                html = reports.generate_html_report(
                    report_data, "http://127.0.0.1:8000"
                ).read_text(encoding="utf-8")
                workbook = load_workbook(
                    reports.generate_excel_report(report_data),
                    read_only=True,
                )
                step_values = list(workbook["Test Steps"].values)
                workbook.close()
            finally:
                reports.REPORTS_DIR = old_reports
            self.assertIn("No screenshot evidence recorded.", html)
            self.assertIn("Unavailable", "\n".join(map(str, step_values[1])))
            self.assertNotIn("outside.png", html)
            self.assertNotIn("outside.png", "\n".join(map(str, step_values[1])))
        finally:
            temporary_directory.cleanup()

    def test_invalid_ai_plan_uses_deterministic_safe_fallback(self):
        FIXTURE_REQUESTS.clear()
        result = self.execute_fixture(
            lambda prompt: '{"actions":[{"action":"submit_form"}]}'
        )
        run_id, run, steps, _, report_data, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertGreaterEqual(len(steps), 3)
            self.assertEqual(steps[0].action, "autonomous:inspect")
            self.assertTrue(any(step.action == "autonomous:fill" for step in steps))
            self.assertTrue(any(step.action == "autonomous:click" for step in steps))
            self.assertFalse(any("password" in step.action for step in steps))
            self.assertFalse(any("/must-not-submit" in (step.page_or_screen or "") for step in steps))
            self.assertFalse(any("/must-not-submit" in url for url in FIXTURE_REQUESTS))
            self.assertTrue(
                any(log.event_type == "planning_fallback" for log in report_data["logs"])
            )
        finally:
            temporary_directory.cleanup()

    def test_planner_service_failure_uses_deterministic_safe_fallback(self):
        def unavailable_planner(prompt: str) -> str:
            raise RuntimeError("Gemini service is unavailable.")

        result = self.execute_fixture(unavailable_planner, max_steps=3)
        _, run, steps, _, report_data, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertEqual(len(steps), 3)
            self.assertEqual(steps[0].action, "autonomous:inspect")
            self.assertTrue(any(step.action == "autonomous:fill" for step in steps))
            self.assertTrue(
                any(log.event_type == "planning_fallback" for log in report_data["logs"])
            )
        finally:
            temporary_directory.cleanup()

    def test_unobserved_metadata_navigation_is_rejected_and_falls_back(self):
        FIXTURE_REQUESTS.clear()
        result = self.execute_fixture(
            lambda prompt: (
                '{"actions":[{"action":"navigate",'
                '"url":"http://169.254.169.254/latest/meta-data/"}]}'
            ),
            max_steps=2,
        )
        _, run, steps, _, report_data, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertEqual(len(steps), 2)
            self.assertEqual(steps[1].action, "autonomous:fill")
            self.assertFalse(
                any("169.254.169.254" in url for url in FIXTURE_REQUESTS)
            )
            self.assertTrue(
                any(log.event_type == "planning_fallback" for log in report_data["logs"])
            )
            self.assertTrue(
                any(log.event_type == "step_limit" for log in report_data["logs"])
            )
        finally:
            temporary_directory.cleanup()

    def test_assertion_failure_is_a_suspected_bug_with_evidence(self):
        result = self.execute_fixture(
            lambda prompt: '{"actions":[{"action":"assert_visible","element_id":"heading-0","expected_text":"Absent heading"}]}'
        )
        run_id, run, steps, bugs, report_data, root, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            failed = next(step for step in steps if step.status == "failed")
            self.assertIn("Assertion failed", failed.error_details)
            self.assertTrue(failed.screenshot_path.endswith("-failure.png"))
            self.assertEqual(len(bugs), 1)
            self.assertIn("Suspected issue", bugs[0].title)
            self.assertEqual(bugs[0].confidence, 0.5)
            evidence_path = root / failed.screenshot_path
            self.assertTrue(evidence_path.is_file())
            with (
                patch.object(reports, "REPORTS_DIR", root / "reports"),
                patch.object(reports, "PROJECT_DIR", root),
            ):
                html = reports.generate_html_report(
                    report_data, "http://127.0.0.1:8000"
                ).read_text(encoding="utf-8")
            self.assertIn(
                f"/test-runs/{run_id}/steps/{failed.step_number}/screenshot",
                html,
            )
        finally:
            temporary_directory.cleanup()

    def test_runtime_limit_and_cancellation_are_recorded(self):
        slow = self.execute_fixture(
            lambda prompt: (time.sleep(1.5) or '{"actions":[{"action":"inspect"}]}'),
            max_duration_seconds=1,
        )
        _, timed_run, timed_steps, _, _, _, temporary_temporary = slow
        try:
            self.assertEqual(timed_run.status, "failed")
            self.assertTrue(any(step.action == "timeout" for step in timed_steps))
        finally:
            temporary_temporary.cleanup()

        cancel_event = threading.Event()

        def cancel_plan(prompt: str) -> str:
            cancel_event.set()
            return '{"actions":[{"action":"fill","element_id":"input-0","value":"query"}]}'

        result = self.execute_fixture(
            cancel_plan,
            max_duration_seconds=15,
            cancel_event=cancel_event,
        )
        _, cancelled_run, cancelled_steps, _, _, _, cancellation_temp = result
        try:
            self.assertEqual(cancelled_run.status, "cancelled")
            self.assertEqual(len(cancelled_steps), 1)
            self.assertEqual(cancelled_steps[0].action, "autonomous:inspect")
        finally:
            cancellation_temp.cleanup()

    def test_navigation_limit_stops_at_five_top_level_navigations(self):
        plans = iter(
            [
                '{"actions":[{"action":"navigate","url":"https://example.com/chain/1"}]}',
                '{"actions":[{"action":"navigate","url":"https://example.com/chain/2"}]}',
                '{"actions":[{"action":"navigate","url":"https://example.com/chain/3"}]}',
                '{"actions":[{"action":"navigate","url":"https://example.com/chain/4"}]}',
                '{"actions":[{"action":"navigate","url":"https://example.com/chain/5"}]}',
            ]
        )
        result = self.execute_fixture(
            lambda prompt: next(plans),
            max_steps=10,
            initial_url="https://example.com/chain/0",
        )
        _, run, steps, _, _, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "failed")
            self.assertTrue(
                any("navigation limit" in (step.error_details or "").lower() for step in steps),
                [(step.action, step.status, step.error_details) for step in steps],
            )
            self.assertLessEqual(len(steps), 6)
        finally:
            temporary_directory.cleanup()

    def test_step_limit_is_enforced_and_reported(self):
        result = self.execute_fixture(
            lambda prompt: '{"actions":[{"action":"inspect"}]}',
            max_steps=2,
        )
        _, run, steps, _, report_data, _, temporary_directory = result
        try:
            self.assertEqual(run.status, "completed")
            self.assertEqual(len(steps), 2)
            self.assertTrue(
                any(log.event_type == "step_limit" for log in report_data["logs"])
            )
        finally:
            temporary_directory.cleanup()


class TestAutonomousRunApi(unittest.TestCase):
    def test_api_creates_and_starts_distinct_exploration_mode(self):
        with TemporaryDirectory() as temporary_directory:
            original_database_path = database.DATABASE_PATH
            database.DATABASE_PATH = Path(temporary_directory) / "api.db"
            try:
                with (
                    patch.dict(os.environ, {"QAI_ALLOWED_HOSTS": ""}),
                    TestClient(app, headers={"Host": "127.0.0.1:8000", "Authorization": f"Bearer {API_TOKEN}"}) as client,
                ):
                    created = client.post(
                        "/test-runs",
                        json={
                            "run_name": "Autonomous API run",
                            "platform": "website",
                            "mode": "autonomous-exploration",
                            "target": "https://example.com/",
                            "goal": "Inspect safely.",
                            "allowed_hosts": ["docs.example.com"],
                            "max_steps": 8,
                            "max_duration_seconds": 30,
                        },
                    )
                    self.assertEqual(created.status_code, 201)
                    run_id = created.json()["id"]
                    self.assertEqual(created.json()["mode"], "autonomous-exploration")
                    self.assertEqual(
                        database.get_test_run(run_id).allowed_hosts,
                        ["docs.example.com"],
                    )

                    def complete_without_browser(run_id, cancel_event, *, planner):
                        self.assertFalse(cancel_event.is_set())
                        self.assertTrue(callable(planner))
                        database.update_test_run_status(run_id, "completed")
                        with website_runner._ACTIVE_RUNS_LOCK:
                            website_runner._ACTIVE_RUNS.pop(run_id, None)
                        website_runner.release_run_slot()

                    with (
                        patch("backend.app.validate_target_url", return_value=FIXTURE_HOST),
                        patch(
                            "backend.app.execute_autonomous_website_run",
                            side_effect=complete_without_browser,
                        ) as execute,
                    ):
                        started = client.post(f"/test-runs/{run_id}/start")

                    self.assertEqual(started.status_code, 202)
                    self.assertEqual(started.json()["mode"], "autonomous-exploration")
                    execute.assert_called_once()
                    history = client.get("/test-runs?platform=website")
                    self.assertEqual(history.status_code, 200)
                    self.assertEqual(history.json()["items"][0]["mode"], "autonomous-exploration")
            finally:
                database.DATABASE_PATH = original_database_path


if __name__ == "__main__":
    unittest.main()
